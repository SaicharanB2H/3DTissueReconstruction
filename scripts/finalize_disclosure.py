"""Finalize the current v2 document without redrawing the approved figures."""
import hashlib
import json
import shutil
import sys
from pathlib import Path
from copy import deepcopy

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.review_tools'))
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = ROOT / 'Invention_Disclosure_Revised_v2.docx'
TARGET = ROOT / 'Invention_Disclosure_Final.docx'
OUT = ROOT / 'results/review/disclosure_final'
REV = ROOT / 'results/review/disclosure_revision'
ROWS = {1:1,2:3,3:5,4:7,5:10,6:12,7:14,8:16,9:18,10:20,11:22,
        12:24,13:26,14:28,15:30,16:32,17:34,18:36,19:38,20:41,21:43,
        22:45,23:47,24:49,25:52,26:54,27:56,28:64,29:66}

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def text_blocks(cell, lines):
    for p in list(cell.paragraphs):
        if not p._p.xpath('.//w:drawing'):
            p._p.getparent().remove(p._p)
    paragraphs=[]
    for line in lines:
        p=cell.add_paragraph(line)
        paragraphs.append(p)
    # Put answer paragraphs before any nested table.
    nested=cell._tc.find(qn('w:tbl'))
    if nested is not None:
        for p in paragraphs:
            nested.addprevious(p._p)
        cell.add_paragraph()
    return paragraphs

def widths(table, sizes):
    table.autofit=False
    table.alignment=WD_TABLE_ALIGNMENT.LEFT
    for node in list(table._tbl.tblPr):
        if node.tag in (qn('w:tblpPr'),qn('w:tblOverlap')):
            table._tbl.tblPr.remove(node)
    for column,size in zip(table.columns,sizes):
        column.width=Inches(size)
    for row in table.rows:
        for c,size in zip(row.cells,sizes):
            c.width=Inches(size)
    table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
    table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),str(round(sum(sizes)*1440)))

def cell_format(cell, size=10.5):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    pr=cell._tc.get_or_add_tcPr()
    margin=pr.find(qn('w:tcMar'))
    if margin is None:
        margin=OxmlElement('w:tcMar');pr.append(margin)
    for edge,value in [('top',65),('bottom',65),('left',90),('right',90)]:
        x=margin.find(qn('w:'+edge))
        if x is None: x=OxmlElement('w:'+edge);margin.append(x)
        x.set(qn('w:w'),str(value));x.set(qn('w:type'),'dxa')
    for p in cell.paragraphs:
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        pf=p.paragraph_format
        pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(0)
        pf.space_before=Pt(0);pf.space_after=Pt(5)
        pf.line_spacing=1.05;pf.widow_control=True
        pf.keep_with_next=False;pf.keep_together=False
        for r in p.runs:
            r.font.name='Calibri';r.font.size=Pt(size)

def image_hashes(doc):
    return [hashlib.sha256(doc.part.related_parts[s._inline.graphic.graphicData.pic.blipFill.blip.embed].blob).hexdigest() for s in doc.inline_shapes]

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    backup=OUT/'Invention_Disclosure_v2_Before_Final.docx'
    if not backup.exists(): shutil.copy2(SOURCE,backup)
    original_sha=sha(SOURCE)
    doc=Document(SOURCE);table=doc.tables[0]
    before_images=image_hashes(doc)
    previous=json.loads((REV/'disclosure_sections.json').read_text(encoding='utf-8'))
    answers={int(k):v for k,v in previous.items()}
    answers={k:v.replace('sample standard deviations across ten volume-level means',
                         'sample standard deviations across ten per-volume metric values') for k,v in answers.items()}
    mismatches=[n for n,r in ROWS.items() if previous[str(n)].strip() not in table.rows[r].cells[1].text.strip()]
    if mismatches:
        raise RuntimeError(f'Reconcile saved user edits before finalization: {mismatches}')

    # Remove repeated full paragraphs; keep the complete evidence in its natural section.
    for n in (14,17,19,23):
        lines=answers[n].split('\n')
        lines=[p for p in lines if not p.startswith('The selected configuration is the corrected original raw 25-step reference;')]
        answers[n]='\n'.join(lines)
    for n in (14,16):
        paragraphs=answers[n].split('\n')
        if n==14:
            paragraphs=[p for p in paragraphs if not p.startswith('Provenance to be supplied:')]
        else:
            paragraphs=[p.split('Provenance to be supplied:')[0].strip() for p in paragraphs]
        answers[n]='\n'.join(paragraphs)+'\nData, conversion and checkpoint provenance limitations are stated in Q17. Experiment results and their limits are stated in Q15.'
    answers[9]='\n'.join(p for p in answers[9].split('\n') if not p.startswith('Earliest conception date'))+'\nConception and public-disclosure chronology: see Q18.'
    for n in (20,22):
        paragraphs=answers[n].split('\n')
        paragraphs=[p.split('Attribution record:')[0].strip() for p in paragraphs]
        answers[n]='\n'.join(p for p in paragraphs if p)+'\nOriginal contributions and later review corrections are distinguished in Q29; individual human attribution requires inventor confirmation.'
    answers[23]=answers[23].replace('Further workable ranges and any earlier supporting experiments Requires inventor confirmation.',
        'Further workable ranges and evidence of any earlier supporting experiments: Requires inventor confirmation.')
    answers[26]=answers[26].replace('confirmation.;','confirmation;')
    answers[17]=answers[17].replace('lineage Requires inventor confirmation.', 'lineage: Requires inventor confirmation.')
    answers[29]=answers[29].replace('Which features each human conceived, and when, Requires inventor confirmation.',
        'Individual human conception contributions and dates: Requires inventor confirmation.')
    answers[29]=answers[29].replace('Revised on 8 October 2026', 'Editorially finalized on 9 October 2026')
    answers[12]=('The approved software diagrams are embedded in two sets, as explicitly requested by this form. '
        'Set A contains the fully labeled architecture and operation flow, shortlisted as the primary explanatory drawings. '
        'Set B contains the same diagrams without descriptions, legends, labels or numbers. Those clean copies are intentional and are not unfinished diagrams. '
        'The drawing designs are unchanged. No physical-device CAD prototype is claimed; historical Blender meshes are visualization artifacts.\n'
        'Drawing index: Figure 1, labeled architecture (101-110); Figure 2, labeled operation flow (201-206); '
        'Figure 3, clean architecture; Figure 4, clean operation flow. Q13 explains the numbered elements.')
    answers[13]=('Figure 1 - architecture: sparse acquired slices and their depth indices (101) enter acquired-only normalization (102) '
        'and the shared residual convolutional context encoder (103). Fourier coordinates (104) describe global depth, relative offset and bounding-slice gap. '
        'Conditioning tokens and spatial maps (105) guide the 2D latent U-Net and DDIM updates (107), starting from seeded Gaussian latent noise (106). '
        'The frozen 2D VAE decoder (108) produces missing planes; assembly and preservation (109) restore normalized acquired values after optional smoothing; '
        'volume export (110) stores the result and provenance. The bypass path carries acquired values directly to assembly.\n'
        'Figure 2 - operation: validate volume, channels and supplied acquired indices (201); normalize using only acquired intensities (202); '
        'find missing depths and bounding acquired contexts (203); encode context and run conditional latent DDIM (204); '
        'decode, optionally smooth and restore acquired slices (205); export and evaluate against withheld ground truth when available (206).\n'
        'Figure 3 is the clean copy of Figure 1. Figure 4 is the clean copy of Figure 2. '
        'The learned models are 2D; volume assembly stacks the planes. Exported geometry remains voxel-index based unless verified physical geometry is supplied.')
    answers[25]=('Specific market size and adoption: a defensible addressable market, installed base or adoption rate for this prototype has not been established. '
        'The table provides industry context and clearly marks its limits.\n'
        'Product category: manufacturer materials document MRI AI reconstruction products, including Siemens Deep Resolve (R8), GE AIR Recon DL (R13), '
        'Canon AiCE (R15), and Fujifilm Synergy DLR (R16). Their availability does not establish demand for this prototype.\n'
        'Customer validation proposed: interview imaging researchers and manufacturers; identify research workflows, permitted data access and integration requirements; '
        'then collect evidence for adoption, pricing and willingness to pay. No project-specific sales forecast, scan-time savings, clinical benefit or economic benefit is asserted.')
    # Lay out the equations separately so each can be read without a dense block.
    lines=answers[15].split('\n')
    expanded=[]
    for line in lines:
        if line.startswith('Coordinates:'):
            for marker in [' Normalization:',' Clean latent:',' Forward diffusion:',' Noise objective:',
                           ' Deterministic DDIM:',' VAE objective:',' Reconstruction MSE=',' SSIM uses']:
                line=line.replace(marker,'\n'+marker.strip())
        expanded.extend(line.split('\n'))
    answers[15]='\n'.join(expanded)
    answers[15]=answers[15].replace('capped at 100 dB for near-perfect matches',
        'with a 100 dB sentinel returned when MSE <= 1e-10')
    answers[15]+=('\nWithin each volume, PSNR is computed from MSE pooled across all missing-plane pixels; SSIM is averaged across 2D slices/channels. '
        'The reported group mean and sample standard deviation summarize the ten per-volume values. '
        'DDIM sampling follows the established implicit diffusion formulation (R20), with eta=0 here. '
        'The VAE posterior sampling noise and the forward diffusion noise are separate Gaussian draws. '
        'Normalization maps intensities <= 1e-7 to background -1; if acquired bounds are degenerate, the implementation uses vmax=vmin+1.')
    # Explicitly name the historical mislabeled baseline in the final disclosure.
    answers[15]+=('\nBaseline terminology: smoothed linear interpolation means linear interpolation followed by Gaussian depth smoothing '
        '(sigma=0.7). The historical Bspline_3d name is retained only as a compatibility alias; it does not denote a B-spline algorithm.')
    for n,row in ROWS.items():
        if n != 12:
            text_blocks(table.rows[row].cells[1],answers[n].split('\n'))
    doc.styles['Normal'].font.name='Calibri'
    doc.styles['Normal'].font.size=Pt(10.5)

    # Reuse the existing embedded picture XML and bytes; do not redraw or retouch.
    figure_cell=table.rows[24].cells[1]
    pictures=[deepcopy(p._p) for p in figure_cell.paragraphs if p._p.xpath('.//w:drawing')]
    for p in list(figure_cell.paragraphs): p._p.getparent().remove(p._p)
    text_blocks(figure_cell,answers[12].split('\n'))
    captions=[
        ('Set A - labeled drawings', 'Figure 1. Architecture: components 101-110',pictures[0]),
        (None,'Figure 2. Operation flow: steps 201-206',pictures[2]),
        ('Set B - clean copies required by Q12','Figure 3. Architecture without labels or numbers',pictures[1]),
        (None,'Figure 4. Operation flow without labels or numbers',pictures[3])]
    for group,caption,picture in captions:
        figure_table=figure_cell.add_table(rows=1, cols=1)
        widths(figure_table,[6.50])
        figure_body=figure_table.cell(0,0)
        figure_body.paragraphs[0].text=(group+'\n' if group else '')+caption
        figure_body._tc.append(picture)
        # Word compatibility mode ignores page-break-before inside the form.
        # Keep each caption and drawing together in an indivisible inline row.
        figure_table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    # The explanatory answer below the nine-column TRL grid needs full width.
    trl_answer=table.rows[64].cells[0].merge(table.rows[64].cells[-1])
    text_blocks(trl_answer,answers[28].split('\n'))
    # The original market table is kept, with accurate headings and balanced columns.
    market=table.rows[52].cells[1].tables[0]
    data=[
      ['Period / scope','Evidence and market-size limits','Growth / adoption context','Relevant products'],
      ['2025 industry context','Philips group sales: EUR 17.8bn (R10). Diagnosis & Treatment segment: EUR 8.5bn (R11). These are company revenues, not this invention\'s market size.','Established MRI AI reconstruction category; no measured prototype adoption.','Deep Resolve (R8); AIR Recon DL (R13); AiCE (R15); Synergy DLR (R16).'],
      ['2026-2028 forecast','Philips forecasts mid-single-digit comparable-sales CAGR for Diagnosis & Treatment (R11). This is a company segment forecast.','Manufacturer forecast; it does not predict project sales or demand.','No commercial product or paying customer established for this project.'],
      ['Prototype opportunity','Specific market size, pricing and projected growth remain unquantified.','Research-only FLAIR reconstruction; independent customer validation proposed.','Potential contacts appear in Q24; interest remains unconfirmed.']]
    for row,values in zip(market.rows,data):
        for c,value in zip(row.cells,values): c.text=value
    widths(market,[0.88,2.30,1.60,1.67])
    widths(table.rows[43].cells[1].tables[0],[0.30,1.35,1.85,2.95])
    table.rows[43].cells[1].tables[0].cell(0,0).text='#'
    widths(table.rows[47].cells[1].tables[0],[1.00,1.80,2.10,1.55])

    # Fit the original 10-column merged form within the A4 text width.
    sec=doc.sections[0]
    usable=sec.page_width-sec.left_margin-sec.right_margin
    grid=table._tbl.tblGrid.gridCol_lst
    old=[int(c.w) for c in grid]
    factor=usable/sum(old)
    new=[round(x*factor) for x in old]
    for c,w in zip(grid,new): c.w=w
    for row in table.rows:
        i=0
        for tc in row._tr.tc_lst:
            span=tc.grid_span
            tc.get_or_add_tcPr().get_or_add_tcW().w=round(sum(new[i:i+span])/635)
            i+=span
    table.autofit=False
    table.alignment=WD_TABLE_ALIGNMENT.LEFT
    table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
    table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),str(round(usable/635)))
    for x in doc.element.xpath('.//w:tblpPr | .//w:tblOverlap'):
        x.getparent().remove(x)
    # Auto row heights permit long answers to continue across pages.
    for x in doc.element.xpath('.//w:trHeight'):
        x.getparent().remove(x)
    for x in doc.element.xpath('.//w:cantSplit'):
        x.getparent().remove(x)
    for c in table._cells:
        cell_format(c)
    question_rows=sorted(set(range(67))-set(ROWS.values())-{8,39,50,57,58,59,61,62,63})
    for index in question_rows:
        for c in table.rows[index].cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next=True
                for run in p.runs: run.font.bold=True
    for row_index in (43,47,52):
        for nested in table.rows[row_index].cells[1].tables:
            for i,row in enumerate(nested.rows):
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                for c in row.cells:
                    cell_format(c,10)
                    if i==0:
                        for p in c.paragraphs:
                            for r in p.runs: r.bold=True
                        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF5');c._tc.get_or_add_tcPr().append(shading)
            if not nested.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None:
                nested.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    # Keep the dense TRL grid legible within the original form structure.
    for row_index in (61,62,63):
        for c in table.rows[row_index].cells:
            cell_format(c,8)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
    for nested in figure_cell.tables:
        for c in nested.rows[0].cells:
            cell_format(c)
            nested.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for p in c.paragraphs:
                if p.text.startswith(('Set A','Set B','Figure ')):
                    p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.bold=True;r.font.size=Pt(11)
                if p._p.xpath('.//w:drawing'):
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_together=True
    for i,s in enumerate(doc.inline_shapes):
        ratio=s.height/s.width
        s.width=Inches(6.35 if i in (0,2) else 5.20)
        s.height=round(s.width*ratio)
        s._inline.docPr.set('descr', ['Approved labeled architecture','Approved labeled operation flow','Clean architecture copy required by Q12','Clean operation flow copy required by Q12'][i])
    # Keep each compact evidence table on one page in Word compatibility mode.
    for row_index in (43,47):
        answer_cell=table.rows[row_index].cells[1]
        evidence_table=answer_cell.tables[0]
        wrapper=answer_cell.add_table(rows=1,cols=1)
        widths(wrapper,[6.55])
        wrapper_body=wrapper.cell(0,0)
        cell_format(wrapper_body,10)
        margins=wrapper_body._tc.get_or_add_tcPr().find(qn('w:tcMar'))
        for edge in margins:edge.set(qn('w:w'),'0')
        wrapper_body.paragraphs[0]._p.getparent().remove(wrapper_body.paragraphs[0]._p)
        evidence_table._tbl.addprevious(wrapper._tbl)
        wrapper_body._tc.append(evidence_table._tbl)
        wrapper_body.add_paragraph().paragraph_format.space_after=Pt(0)
        wrapper.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for n in (24,29):
        for p in table.rows[ROWS[n]].cells[1].paragraphs:
            p.paragraph_format.keep_together=True

    # Break long references onto readable lines and keep source titles with their URLs.
    for p in doc.paragraphs:
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs: r.font.name='Calibri';r.font.size=Pt(10)
        if p.text=='Source register and revision evidence':
            p.text='Source register and supporting evidence'
            p.paragraph_format.page_break_before=True
            p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
        if p.text.startswith('Evidence attachments:'):
            p.text=('Supporting records: results/review/disclosure_final/FINALIZATION_REPORT.md and verification.json; '
                    'results/review/disclosure_revision/sources.json, PRIOR_ART_REVISION.md and PENDING_CONFIRMATIONS.md; '
                    'the experiment CSV/JSON files named in Q15. Remaining personal, ownership, chronology and provenance facts require inventor confirmation.')
        if p.text.startswith('R9.'):
            p.text=p.text.replace('2026 SEC filing','SEC filing dated 19 February 2026')
        if p.text.startswith('R') and p.text.split('.')[0][1:].isdigit():
            p.paragraph_format.keep_together=True
        for r in p.runs:
            r.font.name='Calibri'
            if p.text!='Source register and supporting evidence':r.font.size=Pt(10)
    reference=doc.add_paragraph('R20. Jiaming Song; Chenlin Meng; Stefano Ermon. Denoising Diffusion Implicit Models. '
        'arXiv:2010.02502; v1 6 October 2020; v4 5 October 2022; ICLR 2021. Accessed 8 October 2026.\nhttps://arxiv.org/abs/2010.02502')
    for run in reference.runs:run.font.name='Calibri';run.font.size=Pt(10)
    supporting=next((p for p in doc.paragraphs if p.text.startswith('Supporting records:')),None)
    if supporting is not None:
        supporting._p.addprevious(reference._p)
    footer=sec.footer.paragraphs[0]
    if not footer.text.strip():
        footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run('Invention disclosure | Page ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
        for r in footer.runs: r.font.size=Pt(9)
    doc.core_properties.subject='Final editorial version; unconfirmed inventor and provenance facts explicitly identified'
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.save(TARGET)
    check=Document(TARGET)
    after_images=image_hashes(check)
    alltext='\n'.join(x.text or '' for x in check.element.xpath('.//w:t'))
    result={'source':str(SOURCE),'source_sha256':original_sha,'source_unchanged':sha(SOURCE)==original_sha,
            'final':str(TARGET),'final_sha256':sha(TARGET),'outer_rows':len(check.tables[0].rows),
            'sections':len(ROWS),'embedded_figures':len(check.inline_shapes),
            'diagram_bytes_unchanged':sorted(before_images)==sorted(after_images),
            'floating_tables':len(check.element.xpath('.//w:tblpPr')),
            'title_words':len(answers[1].split()),
            'encoding_issues':[x for x in ['\ufffd','â€™','â€“'] if x in alltext],
            'Q25_column_widths_inches':[round(x.width.inches,3) for x in check.tables[0].rows[52].cells[1].tables[0].columns],
            'native_render':'pending','inventor_confirmations':'Unanswered factual fields remain explicitly marked.'}
    assert result['source_unchanged'] and result['diagram_bytes_unchanged']
    assert result['outer_rows']==67 and result['embedded_figures']==4 and not result['floating_tables']
    assert not result['encoding_issues'] and result['title_words']<=15
    (OUT/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    (OUT/'disclosure_sections.json').write_text(json.dumps(answers,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
