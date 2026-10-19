"""Create the individual report from the retained academic template."""
from pathlib import Path
import sys, copy, csv, json, hashlib, re, textwrap
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.review_tools'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E

OUT = ROOT / 'results/review/project_report'
TEMPLATE = OUT / 'Report_Template_Preserved.docx'
FINAL = ROOT / 'BCSE497J_Project_Report_23BAI1536.docx'
TITLE = '3D MEDICAL VOLUME RECONSTRUCTION VIA LATENT DIFFUSION'
PROGRAM = 'Bachelor of Technology in Computer Science and Engineering with Specialization in Artificial Intelligence and Machine Learning'
doc = Document(TEMPLATE)
body = doc._element.body
for child in list(body): body.remove(child)
sect = OxmlElement('w:sectPr'); body.append(sect)
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin, s.right_margin = Cm(3.81), Cm(2.54)
s.top_margin = s.bottom_margin = Cm(2.54)
s.header_distance = s.footer_distance = Cm(1.27)
for hf in (s.header, s.footer):
    for child in list(hf._element): hf._element.remove(child)
doc.settings.odd_and_even_pages_header_footer = False
s.different_first_page_header_footer = False
for style in doc.styles:
    if style.type == 1:
        style.font.name = 'Times New Roman'; style.font.color.rgb = RGBColor(0,0,0)
normal = doc.styles['Normal']
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name, size in [('Title',20),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    st = doc.styles[name]; st.font.size = Pt(size); st.font.bold = True
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(12); st.paragraph_format.space_after = Pt(8)
    for border in st._element.xpath('.//w:pBdr'): border.getparent().remove(border)
for name in ['TOC 1','TOC 2','TOC 3','Caption']:
    if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    if name in doc.styles:
        doc.styles[name].font.name = 'Times New Roman'
        doc.styles[name].font.size = Pt(12 if name.startswith('TOC') else 11)
        doc.styles[name].paragraph_format.line_spacing = 1
        doc.styles[name].paragraph_format.space_after = Pt(4)
        doc.styles[name].paragraph_format.space_before = Pt(0)
        if name.startswith('TOC'):
            doc.styles[name].font.bold=False
            doc.styles[name].paragraph_format.keep_with_next=False

def field(p, instruction, placeholder=''):
    r = OxmlElement('w:r'); b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin');r.append(b);p._p.append(r)
    r = OxmlElement('w:r'); t = OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=instruction;r.append(t);p._p.append(r)
    r = OxmlElement('w:r');b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'separate');r.append(b);p._p.append(r)
    p.add_run(placeholder)
    r=OxmlElement('w:r');b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'end');r.append(b);p._p.append(r)

p=s.footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(p,' PAGE ', '1')
pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'decimal');pg.set(qn('w:start'),'1');s._sectPr.append(pg)
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings._element.append(upd)

def para(text='', align=None, size=None, bold=False, italic=False, after=6):
    p=doc.add_paragraph();p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after=Pt(after)
    r=p.add_run(text);r.bold=bold;r.italic=italic
    if size:r.font.size=Pt(size)
    return p

def page(): doc.add_page_break()
def heading(text,level=1):
    p=doc.add_paragraph(text,style=f'Heading {level}')
    if level==1:p.paragraph_format.page_break_before=True
    return p
def front(text):
    p=para(text,WD_ALIGN_PARAGRAPH.CENTER,14,True,after=18)
    p.runs[0].underline=True
    p.paragraph_format.keep_with_next=True
    return p
def logo(width=5.0):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT/'template_image2.png'),width=Cm(width))
    p.paragraph_format.space_after=Pt(10)

def cover(first=True):
    start=len(doc.paragraphs)
    if first:para('A project report on',WD_ALIGN_PARAGRAPH.CENTER,12,italic=True,after=10)
    p=doc.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(16)
    para('Submitted in partial fulfillment of the requirements\nfor the award of the degree of',WD_ALIGN_PARAGRAPH.CENTER,14,italic=True,after=8)
    para(PROGRAM,WD_ALIGN_PARAGRAPH.CENTER,22,True,after=10)
    para('by',WD_ALIGN_PARAGRAPH.CENTER,14,italic=True,after=8)
    for name,reg in [('GUDIKANDULA SAI CHARAN CHARY','23BAI1536'),('AYUSH CHAURASIA','23BAI1543'),('CHAHANNA NIROV SHETH','23BAI1505')]:
        para(name,WD_ALIGN_PARAGRAPH.CENTER,16,True,after=0)
        para(reg,WD_ALIGN_PARAGRAPH.CENTER,12,after=6)
    logo()
    para('SCHOOL OF COMPUTER SCIENCE AND ENGINEERING',WD_ALIGN_PARAGRAPH.CENTER,16,True,after=3)
    para('Vellore Institute of Technology\nChennai',WD_ALIGN_PARAGRAPH.CENTER,12,after=8)
    para('October 2026',WD_ALIGN_PARAGRAPH.CENTER,12)
    # Cover roles retain the template's sizes; single spacing keeps each
    # complete cover on one page despite the confirmed long programme name.
    for p in doc.paragraphs[start:]:p.paragraph_format.line_spacing=1

cover();page();cover(False);page()
logo();front('DECLARATION')
para(f'I, Gudikandula Sai Charan Chary (23BAI1536), declare that the project report entitled “{TITLE.title()}”, submitted to Vellore Institute of Technology, Chennai, in partial fulfillment of the requirements for the award of the degree of {PROGRAM}, presents the project work carried out under the guidance of Dr. Sarita Singh, Assistant Professor, School of Computer Science and Engineering.')
para('The methods, implementation, experimental records and limitations described in this report are supported by the project repository and cited sources. Material drawn from published work has been acknowledged.')
para('Place: Chennai\nDate: ____________________',WD_ALIGN_PARAGRAPH.LEFT,after=24)
para('Signature of the student: ____________________\nGudikandula Sai Charan Chary\n23BAI1536',WD_ALIGN_PARAGRAPH.RIGHT)
page();logo();para('SCHOOL OF COMPUTER SCIENCE AND ENGINEERING',WD_ALIGN_PARAGRAPH.CENTER,14,True);front('CERTIFICATE')
para(f'This is to certify that the project report entitled “{TITLE.title()}”, submitted by Gudikandula Sai Charan Chary (23BAI1536), in partial fulfillment of the requirements for the award of the degree of {PROGRAM} at Vellore Institute of Technology, Chennai, is a record of the project work presented under the guidance of Dr. Sarita Singh, Assistant Professor, School of Computer Science and Engineering.')
para('Guide\nDr. Sarita Singh\nAssistant Professor\nSignature: ____________________',WD_ALIGN_PARAGRAPH.LEFT,after=15)
para('Head of Department\nDr. Tulasi Prasad Sarika\nSignature: ____________________',WD_ALIGN_PARAGRAPH.LEFT,after=15)
para('Internal examiner: ____________________\nExternal examiner: ____________________\nDate: ____________________\nInstitution seal: ____________________',WD_ALIGN_PARAGRAPH.LEFT)
page();front('ABSTRACT')
para('This project investigates retrospective reconstruction of a dense brain MRI volume from a sparse set of axial slices. A shared convolutional encoder extracts information from the two acquired slices that bound each missing depth. Fourier coordinate embeddings represent the requested depth, its relative location and the acquired gap. A conditional two dimensional U-Net denoises a compact latent representation, and a frozen variational autoencoder decodes the generated slice. The slices are assembled into a three dimensional volume. Acquired planes are restored exactly in the normalised input space after optional depth smoothing.')
para('The review corrected acquired-only intensity normalisation, missing-target sampling, handling of irregular acquired indices, checkpoint validation, resume and exponential moving average state, and research volume exports. The audited local collection contains 369 complete volumes with 155 slices each. The implemented channel convention assumes FLAIR; the original HDF5 conversion provenance and physical geometry require further confirmation.')
para('A bounded local experiment compared a corrected reference with a fine-tuned model on five validation volumes and ten held-out benchmark volumes at strides 8 and 16. Fine-tuning completed 3,000 steps with equally sampled training strides 4, 8 and 16. The full bounded protocol used 40.36 minutes of the two-hour allocation. The candidate did not satisfy the validation criterion, so the corrected original checkpoint was retained with raw weights, 25 DDIM steps and smoothing. Mean missing-slice PSNR was 28.037 dB at stride 8 and 25.071 dB at stride 16; corresponding SSIM values were 0.9054 and 0.8818. Linear interpolation had higher SSIM at stride 8, while diffusion had higher PSNR at both strides. The results establish a reproducible research prototype, with no demonstrated clinical efficacy, segmentation improvement or scanner acquisition-time benefit.')
para('Keywords: sparse MRI, slice reconstruction, latent diffusion, variational autoencoder, DDIM, acquired-slice preservation.',italic=True)
page();front('ACKNOWLEDGEMENT')
para('I express my gratitude to my project guide, Dr. Sarita Singh, Assistant Professor, School of Computer Science and Engineering, Vellore Institute of Technology, Chennai, for guidance during this project. I thank Dr. Tulasi Prasad Sarika, Head of Department, and the faculty of the school for the academic environment in which this work was developed.')
para('I acknowledge Vellore Institute of Technology, Chennai, for the opportunity to undertake this project as part of the Bachelor of Technology programme in Computer Science and Engineering with Specialization in Artificial Intelligence and Machine Learning. I also acknowledge the researchers whose publications form the foundation of the methods discussed in this report, and the BraTS data community for the source dataset context.')
para('I thank my family and those who supported my learning and completion of the project. This individual report presents the implemented system, the corrected experimental evidence and the remaining limitations.')
para('Gudikandula Sai Charan Chary\n23BAI1536',WD_ALIGN_PARAGRAPH.RIGHT,after=0)
page();front('TABLE OF CONTENTS');field(doc.add_paragraph(),' TOC \\o "1-2" \\h \\z \\u ','Contents will be updated in Word.')
page();front('LIST OF TABLES');table_list_anchor=doc.paragraphs[-1]
page();front('LIST OF FIGURES');figure_list_anchor=doc.paragraphs[-1]
page();front('SYMBOLS ABBREVIATIONS AND NOMENCLATURE')

book_id=1;tables=[];figures=[]
def bookmark(p,name):
    global book_id
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(book_id));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(book_id));book_id+=1
    p._p.insert(0,a);p._p.append(b)
def caption(label,title,records):
    name=('tbl' if label=='Table' else 'fig')+str(len(records)+1)
    text=f'{label} {len(records)+1}  {title}'
    p=para(text,WD_ALIGN_PARAGRAPH.CENTER,11,after=6)
    p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
    bookmark(p,name);records.append((text,name));return p
def table(title,headers,rows,widths=None):
    caption('Table',title,tables)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [14.65/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for k,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for tag,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                x=OxmlElement('w:'+tag);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');mar.append(x)
            pr.append(mar)
            if k==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.size=Pt(12);r.bold=k==0
    para('',after=3)
    return t
def figure(path,title,width=14.4):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(ROOT/path),width=Cm(width))
    c=caption('Figure',title,figures);c.paragraph_format.keep_with_next=False

table('Abbreviations and symbols',['Term','Meaning'],[
('MRI','Magnetic resonance imaging'),('FLAIR','Fluid attenuated inversion recovery'),('VAE','Variational autoencoder'),('U-Net','Encoder and decoder neural network with skip connections'),('DDPM','Denoising diffusion probabilistic model'),('DDIM','Denoising diffusion implicit model'),('EMA','Exponential moving average'),('PSNR','Peak signal to noise ratio'),('SSIM','Structural similarity index measure'),('MSE and MAE','Mean squared error and mean absolute error'),('AMP','Automatic mixed precision'),('D C H W','Depth, channels, height and width'),('A and M','Acquired and missing depth index sets'),('z and t','Slice depth index and diffusion timestep'),('α and β','Relative interpolation position and diffusion variance schedule'),('θ','Trainable model parameters')],[3.4,11.25])

# Native Word equations authored through Microsoft's MathML to OMML stylesheet.
transform=E.XSLT(E.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
def mi(x):return '<mi>'+x+'</mi>'
def mn(x):return '<mn>'+str(x)+'</mn>'
def mo(x):return '<mo>'+x+'</mo>'
def row(*xs):return '<mrow>'+''.join(xs)+'</mrow>'
def sub(x,y):return '<msub>'+mi(x)+mi(y)+'</msub>'
def sup(x,y):return '<msup>'+x+mn(y)+'</msup>'
def frac(a,b):return '<mfrac>'+a+b+'</mfrac>'
def sqrt(a):return '<msqrt>'+a+'</msqrt>'
def bar(x):return '<mover accent="true">'+mi(x)+mo('¯')+'</mover>'
def abar(t):return '<msub>'+bar('α')+mi(t)+'</msub>'
def equation(xml,label):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    result=transform(E.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+xml+'</math>').encode()))
    m=copy.deepcopy(result.getroot())
    for r in m.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}r'):
        pr=E.Element(qn('w:rPr'));fonts=E.SubElement(pr,qn('w:rFonts'));fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math')
        sz=E.SubElement(pr,qn('w:sz'));sz.set(qn('w:val'),'24');r.insert(0,pr)
    p._p.append(m);p.add_run('  ('+label+')');return p

def prose(block):
    for text in block.strip().split('\n\n'):para(' '.join(text.splitlines()))

heading('CHAPTER 1 INTRODUCTION')
heading('1.1 PROJECT CONTEXT',2)
prose('''The project estimates missing axial planes of a brain MRI volume from a selected set of acquired planes. Its deliverable is a research reconstruction pipeline, reproducible evaluation records and documentation of the evidence. The current experiment supports a narrower conclusion than the original project proposal: conditional latent diffusion can achieve higher missing-slice PSNR than linear interpolation on the saved benchmark, but it does not improve every fidelity measure. Fine-tuning during the review did not produce an eligible replacement for the corrected original model.

Retrospective sparse reconstruction starts with an already available dense volume and withholds slices to simulate an acquisition pattern. This permits comparison with the withheld ground truth. It does not demonstrate how an MRI scanner would acquire the sparse stack, how much scan time would be saved, or whether a clinician would make the same decision from the reconstruction. Those questions require prospective acquisition studies and clinical evaluation beyond this implementation.

The title refers to the assembled three dimensional output. The implemented encoder, autoencoder and denoising network operate on two dimensional planes. Their conditioning uses neighbouring acquired planes and depth coordinates; they do not perform full-volume three dimensional convolution. This distinction matters when interpreting continuity between independently generated slices.''')
heading('1.2 PROBLEM STATEMENT',2)
prose('''Given the depth of a volume, its channels and a set of acquired slice indices, reconstruct all missing planes while retaining the supplied acquired values exactly. Every requested missing depth must be bounded by two acquired planes. The system must use only available acquisition information for its preprocessing and conditioning. Dense ground truth may be used for supervised training and evaluation, but must not supply normalisation statistics at inference.

Interpolation provides a necessary reference because it directly estimates intermediate intensities without learned weights. It can perform strongly for neighbouring, registered planes and is inexpensive. The learned model is justified only by demonstrated differences on held-out volumes and an explicit account of its failure modes. Perceptually plausible detail can be incorrect, especially for unusual lesions or large gaps. Visual appearance alone is therefore insufficient evidence of anatomical accuracy.''')
heading('1.3 OBJECTIVES',2)
table('Project objectives and acceptance evidence',['Objective','Acceptance evidence'],[
('Reliable preprocessing','Acquired-only percentile bounds; consistent context and target scaling'),('Valid sparse reconstruction','Actual acquired-index complement; boundary and shape checks'),('Data fidelity','Exact restoration of acquired normalised planes'),('Reproducible learning','Explicit initialisation, resume, seed, training stride and run caps'),('Measured comparison','Five validation cases and ten benchmark cases at strides 8 and 16'),('Honest selection','Higher validation PSNR with no decrease in validation SSIM'),('Portable research output','NIfTI and HDF5 layouts with geometry limitations recorded')],[4.6,10.05])
heading('1.4 SCOPE AND CONTRIBUTIONS',2)
prose('''The demonstrated scope is a single channel from a locally converted BraTS brain MRI collection, assumed to be FLAIR by the loader convention. The review retained the custom architecture and the existing VAE and diffusion checkpoints. Work focused on correcting leakage and evaluation behaviour, exposing bounded training controls, and attempting one fine-tuning candidate on local compute. The original source documents and historical results were preserved.

The technical contribution of this project is the implemented combination of shared context encoding, explicit depth and gap conditioning, latent denoising, acquired-plane restoration and an auditable evaluation workflow. These engineering contributions should not be read as an assertion that sparse reconstruction or latent diffusion is a new research concept. Closely related published work exists, and no head-to-head experiment with those systems was carried out.

Clinical efficacy, segmentation Dice improvement, computed tomography, other organs, multicontrast joint reconstruction, scanner integration and acquisition-time savings remain proposed extensions. The software supports multiple channel export shapes, but that does not establish the availability of trained multichannel weights or clinical interoperability.''')
heading('1.5 REPORT ORGANISATION',2)
para('Chapter 2 introduces the relevant methods and literature. Chapter 3 records the data and requirements. Chapter 4 describes the architecture and equations. Chapter 5 explains implementation and corrections. Chapter 6 defines the experiment. Chapter 7 presents the measured results and limitations. Chapter 8 concludes the work and identifies future studies. The appendices provide reproducible commands, detailed tables and source listings for the critical data and reconstruction paths.')

heading('CHAPTER 2 BACKGROUND AND LITERATURE REVIEW')
heading('2.1 MRI DATA AND SPARSE SLICE RECONSTRUCTION',2)
prose('''The official BraTS 2020 data description identifies multimodal brain MRI and associated tumour annotations (Center for Biomedical Image Computing and Analytics [CBICA], 2020). The local project uses converted HDF5 files rather than the source NIfTI files. Consequently, the official source description establishes the dataset context but does not verify the conversion's channel ordering, mask-channel mapping or preserved geometry. The report treats these as provenance limitations.

Sparse slice interpolation differs from k-space reconstruction. In this project the available inputs are image planes at known integer depths. There is no implementation of MRI coil sensitivity estimation, Fourier measurement consistency or a k-space acquisition model. Missing-depth reconstruction therefore cannot be equated with accelerated MRI reconstruction methods that operate on scanner measurements.

Nearest-neighbour interpolation repeats the closest acquired plane. Linear interpolation blends the two surrounding planes according to relative depth. Cubic interpolation uses a higher-order curve where sufficient samples exist. The project additionally includes linear interpolation followed by Gaussian smoothing along depth. Its corrected name is smoothed linear interpolation; a historical B-spline alias is retained only for compatibility.''')
heading('2.2 VARIATIONAL AUTOENCODERS',2)
prose('''Kingma and Welling (2014) introduced a tractable variational learning framework using a reparameterised latent distribution. An encoder predicts a distribution over a latent representation, while a decoder reconstructs the input. The reconstruction term and the regularisation of the latent distribution serve different purposes. The reparameterisation allows gradients to pass through a stochastic latent sample during training.

In this implementation, the custom medical autoencoder compresses a 240 by 240 single-channel plane into four 30 by 30 latent channels. The frozen decoder determines which structures the diffusion model can reproduce. If the autoencoder loses a fine boundary, improving the denoiser alone cannot guarantee its recovery. Keeping the VAE frozen also separates the bounded experiment from a full retraining study, but preserves any mismatch caused by its historical preprocessing.''')
heading('2.3 DIFFUSION MODELS AND DDIM',2)
prose('''Ho et al. (2020) developed denoising diffusion probabilistic models in which training learns to reverse progressive corruption by noise. The project follows the noise-prediction formulation: a clean latent is perturbed at a sampled timestep and the denoiser predicts the added Gaussian noise. Supervised image targets enter through their encoded latent representations, while conditioning describes the available neighbouring slices.

Song et al. (2021) proposed denoising diffusion implicit models, allowing a non-Markovian sampling construction and accelerated sampling using fewer steps. The project uses the implemented DDIM update with zero sampling stochasticity. A seed is still necessary because the initial latent is random. Thus the same settings and seed can repeat a reconstruction in the verified environment, while a different seed may yield different generated anatomy.

Nichol and Dhariwal (2021) studied improvements to diffusion modelling, including a cosine noise schedule. The implemented default uses 1,000 training timesteps with a cosine schedule. This is the training schedule length, not the number of denoising calls at reconstruction. The selected reference performs 25 inference steps, and 50 steps were compared on validation data.''')
heading('2.4 LATENT DIFFUSION AND CONDITIONING',2)
prose('''Rombach et al. (2022) demonstrated image synthesis in an autoencoder latent space and the use of cross-attention for conditioning. This reduces the spatial size of the denoising problem. The current project uses these general ideas with a custom medical VAE and slice-context encoder; it does not use the published model's externally downloaded image or text weights.

Ronneberger et al. (2015) described U-Net's contracting and expanding paths with skip connections for biomedical image segmentation. Here a related encoder-decoder topology predicts latent noise rather than class labels. Calling the network a U-Net therefore does not imply that the project implements a segmentation system. No segmentation model was trained or evaluated in the reported experiment.

The conditioning path has both spatial and token representations. Spatial concatenation supplies local maps at the latent resolution. Cross-attention supplies a sequence of pooled context tokens and a coordinate token. Depth and gap information allow a single denoiser to receive different requested positions, rather than assuming that every missing slice is halfway between its neighbours.''')
heading('2.5 CLOSELY RELATED MEDICAL RECONSTRUCTION WORK',2)
prose('''Bubeck et al. (2025, revised 2026) proposed cardiac latent interpolation diffusion for reconstructing volumes from sparse cardiac MRI slices. Their work overlaps with this project in sparse-slice input, latent diffusion and volumetric output. It also discusses cardiac and temporal data that are outside the current project's demonstrated scope. The existence of this work prevents a broad claim that sparse medical volume reconstruction using latent diffusion alone is unique.

The project-specific distinctions to examine are its explicit global depth, relative depth and normalised gap inputs, hybrid spatial and token conditioning, exact preservation of acquired values, acquired-only normalisation and handling of supplied irregular indices. These are implemented features. Their research novelty, patentability or superiority over another method is not established by the present benchmark. A fair related-method comparison would require the same data, splits, acquisition pattern, preprocessing, compute accounting and evaluation masks.''')
table('Literature relationship to this implementation',['Work','Relevant concept','Project relationship'],[
('Kingma and Welling 2014','Variational latent encoding','Custom VAE; frozen during fine-tuning'),('Ho et al. 2020','Noise-prediction diffusion','Latent noise MSE training'),('Song et al. 2021','DDIM sampling','25 and 50 steps; zero stochasticity'),('Nichol and Dhariwal 2021','Cosine schedule','Default 1,000-step training schedule'),('Rombach et al. 2022','Latent diffusion and conditioning','Custom medical latent and context paths'),('Bubeck et al. 2025','Sparse cardiac volume reconstruction','Related work; no direct comparison')],[4.1,4.45,6.1])
heading('2.6 REVIEW CONCLUSION',2)
para('The literature supports the method family and identifies close prior work. It does not supply evidence for this implementation’s measured results. Those results are established by the saved local experiments described in Chapters 6 and 7. Evaluation integrity, geometry provenance and clinical validation remain separate requirements, even when the model architecture follows established research.')

heading('CHAPTER 3 REQUIREMENTS AND DATA ANALYSIS')
heading('3.1 DATA ORGANISATION',2)
prose('''The local collection contains 57,195 slice files, grouped into 369 complete volumes. Each volume has 155 axial planes of size 240 by 240. The filenames encode a local volume identifier and a slice index. The loader expects an image dataset with height, width and four channel axes, and a mask dataset with height, width and three channels. The reconstruction tensor is reordered into depth, channels, height and width.

The code assumes the four image channels correspond to FLAIR, T1, contrast-enhanced T1 and T2. This assumption was not verified against the original HDF5 converter. Sampled mask channels are mutually exclusive, so they must not be interpreted automatically as the overlapping whole tumour, tumour core and enhancing tumour regions. Masks do not condition the reported reconstruction experiment.

All filenames, depth indices and available metadata were audited; full numerical payloads were sampled rather than every voxel in the collection. The HDF5 files do not supply verified physical spacing, orientation or a patient affine. Local volume identifiers serve as grouping keys but do not independently verify original clinical patient identities. These limits should remain visible in any subsequent export or dataset reuse.''')
table('Data and partition summary',['Item','Recorded value'],[
('Slice files','57,195'),('Complete local volumes','369'),('Volume dimensions','155 × 240 × 240; selected channel count 1'),('Stored image shape','240 × 240 × 4'),('Stored mask shape','240 × 240 × 3'),('Training validation test volumes','258 / 55 / 56, split seed 42'),('Fixed validation volume IDs','2, 5, 9, 15, 28'),('Final benchmark volume IDs','14, 21, 22, 35, 49, 50, 51, 53, 55, 59')],[5.0,9.65])
heading('3.2 DATA SEPARATION',2)
prose('''The split groups slices by local volume ID before assigning training, validation and test partitions. This prevents neighbouring slices of the same local volume from appearing in different partitions. The configured counts are 258, 55 and 56. The final comparison uses only ten cases already represented in the historical benchmark, not the entire internal test partition and not the official BraTS challenge test set.

Five fixed validation cases determine the inference setting and whether the fine-tuned candidate is eligible. The final ten cases are reserved for reporting the selected result. Repeatedly choosing settings based on these test results would turn them into a validation set. This experiment uses the saved selection rule and does not use the test metrics to select a new checkpoint.

Patient separation in this report means separation by local volume identifiers. Stronger assurance requires the original patient mapping and inspection for duplicates or repeated acquisitions. The available evidence supports the internal grouping check; it does not establish independence across institutions or scanners.''')
heading('3.3 FUNCTIONAL REQUIREMENTS',2)
prose('''The reconstruction API must validate the tensor rank, finite intensities, channels, spatial divisibility and acquired indices before generation. Acquired indices must be unique integers within the volume range and include both boundaries. The missing set is computed from their complement. A stride is useful for constructing a retrospective acquisition pattern, but is insufficient to describe a genuinely supplied irregular pattern.

Model creation must follow a single configuration-aware factory. Required model states must be present and compatible with the constructed modules. Missing checkpoints must cause an explicit failure rather than allowing random initial weights to produce plausible-looking outputs. Raw and averaged weights must be selected deliberately and unavailable averaged states must be reported.

Exports must distinguish array layout from physical geometry. A normalised research tensor is not a raw scanner-intensity image, and identity voxel coordinates are not a verified patient coordinate system. The output sidecar records settings, normalisation, checkpoint identity and the limits of available metadata.''')
heading('3.4 NONFUNCTIONAL REQUIREMENTS',2)
table('Operational requirements',['Requirement','Implementation or limitation'],[
('Bounded local compute','Two-hour protocol cap; no cloud training'),('Reproducibility','Fixed split and reconstruction seeds; hashes and manifests'),('Memory control','Batch 8, AMP and compact latent representation'),('Compatibility','Historical entry points and baseline alias retained'),('Auditability','Per-volume CSVs, JSON snapshots and preserved checkpoints'),('Failure visibility','Clear errors for shapes, indices, settings and weights'),('Clinical safety scope','Research use; generated anatomy can be incorrect')],[4.5,10.15])
heading('3.5 DEVELOPMENT ENVIRONMENT',2)
para('The recorded experiment used Python 3.10.10, PyTorch 2.9.1 with CUDA 12.8 and an NVIDIA RTX 5060 Laptop GPU with approximately 8 GB VRAM. The user-provided environment is D:\\Kaggle competetion\\torch_env. CPU operation is exposed by the software but was not benchmarked. The review and report use existing local records; the report creation does not add another training run.')
heading('3.6 ACQUISITION PATTERNS',2)
table('Boundary-inclusive acquisition at depth 155',['Stride','Acquired planes','Missing planes'],[(2,78,77),(4,40,115),(8,21,134),(16,11,144)],[3.65,5.5,5.5])
para('The final boundary at index 154 is included even when the regular stride does not land on it. Consequently, a stride 16 pattern is not exactly one sixteenth of 155 planes. The final gap may differ from the interior gaps. Counting actual acquired planes is more accurate than calling the simulation a sixteen-fold scanner acceleration.')

heading('CHAPTER 4 SYSTEM DESIGN AND METHODOLOGY')
heading('4.1 ARCHITECTURE OVERVIEW',2)
figure(Path('results/review/drawings/architecture_labeled.png'),'Implemented conditional latent diffusion architecture')
para('Figure 1 reuses the approved architecture drawing. The numbered labels identify processing blocks, not additional network layers. The acquired-only preprocessing path supplies both context generation and the preserved planes used during assembly. The latent noise and the conditioning representations enter the U-Net; its final latent is decoded by the frozen VAE. Assembly restores acquired planes before export.')
heading('4.2 ACQUIRED AND MISSING INDEX SETS',2)
equation(row(mi('M'),mo('='),mi('{0,…,D−1}'),mo('∖'),mi('A')),'4.1')
para('A denotes the supplied acquired depths and M the missing depths. For a target depth z, the previous and next acquired indices are its nearest bounding entries in the sorted acquired list. The API requires acquisition at the first and last depth, avoiding unsupported extrapolation. Irregular internal gaps are valid because neighbours are found from the actual list.')
heading('4.3 ACQUIRED ONLY NORMALISATION',2)
equation(row(sub('I','norm'),mo('='),mn(2),mo('·'),frac(row(mi('clip'),mo('('),mi('I'),mo(','),sub('v','min'),mo(','),sub('v','max'),mo(')'),mo('−'),sub('v','min')),row(sub('v','max'),mo('−'),sub('v','min'),mo('+'),mi('ε'))),mo('−'),mn(1)),'4.2')
prose('''The lower and upper bounds are the 0.5th and 99.5th percentiles of foreground values in acquired slices only. Foreground is defined by values greater than the small epsilon threshold. The threshold is one ten-millionth. When the acquired foreground is absent, bounds are derived from the acquired array; degenerate bounds are separated by one intensity unit. All intensities and bounds must be finite.

The same bounds are used for acquired context planes and withheld targets during training and evaluation. Background values at or below epsilon map to negative one. Percentile clipping makes this transform only approximately invertible. Exact acquired preservation therefore refers to the normalised tensor supplied to reconstruction, not to lossless recovery of original scanner intensities.

Estimating percentiles from the complete dense volume would reveal withheld intensities to the system. It would also make real sparse inference depend on information that is unavailable. The acquired-only correction addresses this mismatch and is the reason historical scores with different preprocessing remain separate.''')
heading('4.4 SHARED CONTEXT ENCODER',2)
prose('''The two bounding planes pass through a shared residual convolutional encoder. The configured base width is 32, with stage widths 32, 64, 128 and 128. Three spatial reductions give a 30 by 30 feature grid from 240 by 240 inputs. Group normalisation and SiLU activation are used. This is a custom encoder, not a pretrained ResNet18.

For spatial conditioning, the previous and next 128-channel feature maps are concatenated with a relative-position plane. The resulting 257 channels are fused and projected into eight channels at the latent resolution. For token conditioning, a relative-position blend of the feature maps is adaptively pooled to a four by four grid and projected to 256 dimensions, yielding sixteen spatial tokens. A coordinate token adds one further token, making seventeen tokens in total.''')
equation(row(mi('α'),mo('='),frac(row(mi('z'),mo('−'),sub('z','prev')),row(sub('z','next'),mo('−'),sub('z','prev')))),'4.3')
equation(row(mi('d'),mo('='),frac(mi('z'),row(mi('D'),mo('−'),mn(1))),mo(','),mi('g'),mo('='),frac(row(sub('z','next'),mo('−'),sub('z','prev')),mn(16))),'4.4')
para('The three coordinate inputs are relative position α, global normalised depth d and gap g. The code uses deterministic sinusoidal Fourier features and learned projections, not random Fourier features. Each coordinate produces a 64-dimensional embedding; their concatenation is projected from 192 to 256 dimensions. In the worked example of depth 12 bounded by indices 8 and 16, the relative position and normalised gap are both one half, while global depth is 12 divided by 154.')
heading('4.5 VARIATIONAL AUTOENCODER',2)
table('Network dimensions for the demonstrated input',['Component','Configured representation'],[
('Selected image plane','1 × 240 × 240'),('VAE latent','4 × 30 × 30'),('Spatial context','8 × 30 × 30'),('Context token sequence','17 × 256'),('VAE base and multipliers','32; 1, 2, 4, 4'),('U-Net base and multipliers','64; 1, 2, 4'),('Attention','4 heads; context dimension 256')],[5.1,9.55])
equation(row(sub('z','0'),mo('='),mn('0.18215'),mo('·'),mo('('),mi('μ'),mo('+'),mi('σ'),mo('·'),sub('ε','vae'),mo(')')),'4.5')
para('The VAE posterior is a diagonal Gaussian parameterised by mean and log variance. A reparameterised sample is multiplied by the latent scale 0.18215. The latent spatial dimensions are eight times smaller along each axis. Accounting for four latent channels, a single-channel 240 by 240 image has sixteen times as many scalar elements as its latent. This count is not a measured memory or runtime speedup.')
equation(row(sub('L','VAE'),mo('='),sub('L','1'),mo('+'),mn('0.5'),mi('MSE'),mo('+'),sup(mn(10),'−5'),mi('KL'),mo('+'),mn('0.1'),sub('L','gradient')),'4.6')
para('The recorded VAE objective combines absolute reconstruction error, half the squared error, KL regularisation and a spatial-gradient absolute error. The gradient term uses horizontal and vertical differences. The best VAE checkpoint is stored at epoch 13 with validation objective 0.0891816. This scalar is a training objective, not PSNR. During the review experiment the VAE is frozen and receives no parameter updates.')
heading('4.6 CONDITIONAL LATENT DENOISING',2)
equation(row(sub('z','t'),mo('='),sqrt(abar('t')),mo('·'),sub('z','0'),mo('+'),sqrt(row(mn(1),mo('−'),abar('t'))),mo('·'),mi('ε')),'4.7')
para('In Equation 4.7 the barred alpha at timestep t denotes cumulative retained signal from the diffusion schedule. It is distinct from the relative slice position in Equation 4.3. Forward-process Gaussian noise is independent of the Gaussian draw used for the VAE posterior. A timestep is sampled during training, and the U-Net predicts the forward noise from the noisy latent, timestep and slice context.')
equation(row(sub('L','diffusion'),mo('='),mi('E'),mo('['),sup(row(mo('‖'),mi('ε'),mo('−'),sub('ε','θ'),mo('('),sub('z','t'),mo(','),mi('t'),mo(','),mi('c'),mo(')'),mo('‖')),2),mo(']')),'4.8')
para('The expectation in Equation 4.8 is implemented as mean squared noise error over sampled batches, timesteps and latent elements. The training target is a missing plane. Context c contains spatial maps and tokens. No adversarial, segmentation or explicit three dimensional consistency loss is included in the fine-tuning objective. The U-Net uses residual blocks, skip connections and cross-attention; size alignment is required because the latent grid reduces through 30, 15 and 8 pixels.')
heading('4.7 DDIM RECONSTRUCTION',2)
equation(row(sub('z','clean'),mo('='),frac(row(sub('z','t'),mo('−'),sqrt(row(mn(1),mo('−'),abar('t'))),mo('·'),sub('ε','θ')),sqrt(abar('t')))),'4.9')
equation(row(sub('z','prev'),mo('='),sqrt(abar('prev')),mo('·'),sub('z','clean'),mo('+'),sqrt(row(mn(1),mo('−'),abar('prev'))),mo('·'),sub('ε','θ')),'4.10')
prose('''Equations 4.9 and 4.10 describe the deterministic update used when DDIM stochasticity is zero. Here prev is the preceding inference timestep, not the previous acquired depth. The timestep schedule is generated from the training schedule and the requested inference step count. Inference starts from seeded latent Gaussian noise and repeatedly applies noise prediction and the scheduler update before decoding.

The selected setting uses 25 inference steps. Fifty steps were available and evaluated but did not produce the best validation score in this bounded comparison. More steps are therefore not assumed to improve this checkpoint. The initial random latent and the exact sequence of missing-slice batches are part of reproducibility.''')
heading('4.8 ASSEMBLY AND PROCESS FLOW',2)
figure(Path('results/review/drawings/operation_labeled.png'),'Reconstruction operation flow',width=12.1)
para('The process diagram is reused unchanged. Optional Gaussian smoothing is applied along depth with standard deviation 0.5 for diffusion output. The supplied acquired values are restored afterwards, making fidelity to those input planes an explicit invariant. Smoothing can reduce depth discontinuity, but does not establish consistent anatomy or a calibrated uncertainty estimate. Export and metrics are separate steps; metrics require withheld ground truth.')

heading('CHAPTER 5 IMPLEMENTATION AND RELIABILITY REVIEW')
heading('5.1 SOFTWARE ORGANISATION',2)
table('Implementation module map',['Area','Principal modules'],[
('Data','src/data/preprocessing.py; sparsifier.py; transforms.py'),('Construction','src/models/factory.py'),('Conditioning','slice_encoder.py; positional_embedding.py'),('Generation','vae.py; unet_diffusion.py; scheduler.py; pipeline.py'),('Evaluation','src/evaluation/metrics.py and baseline modules'),('Export','src/utils/volume_io.py'),('Interfaces','scripts/reconstruct_volume.py; evaluate_all.py; finetune_diffusion.py'),('Bounded protocol','scripts/review_experiment.py')],[3.8,10.85])
prose('''A central factory constructs modules from supported settings and loads trained states consistently across training, evaluation and reconstruction. The pipeline then coordinates preprocessing-independent normalised inputs, neighbour selection, encoding, denoising, decoding and assembly. Keeping these responsibilities explicit makes it possible to distinguish a model failure from invalid data or an unsupported configuration.

The project preserves older script names as compatibility entry points. Historical configuration fields do not automatically control every new interface. For example, bounded diffusion fine-tuning uses CLI learning rate, batch size, stride, seed and run caps, while architecture, schedule and supported optimiser settings come from configuration. Unsupported model backbones, prediction types or conditioning modes fail explicitly.''')
heading('5.2 CORRECTED PREPROCESSING AND TARGET SAMPLING',2)
prose('''The original evaluation could derive normalisation from dense information or use inconsistent statistics for different roles. The reviewed path caches acquired-only statistics per volume and acquisition pattern and applies them to both context and targets. Image caching and statistic caching are separate so that repeated training access does not cause redundant HDF5 loading without silently reusing a different acquisition's bounds.

Diffusion training now chooses target depths only from the missing set. Otherwise acquired planes could be used as targets even though the objective is to reconstruct unavailable planes. The VAE remains a separate image reconstruction model, so its training is not restricted in the same way. Equally sampled training strides refer to the stride draw probability, not equal observed counts in every finite run.

The normalisation-independence check changes missing ground-truth intensities while holding acquired slices fixed. The acquired statistics must remain unchanged. This is a targeted leakage test: it establishes the intended dependency of preprocessing, rather than claiming that every possible source of data leakage has been eliminated.''')
heading('5.3 INDEX VALIDATION AND ACQUIRED FIDELITY',2)
prose('''The missing set is determined from the supplied indices rather than inferred solely from a nominal stride. Unsorted valid acquired lists are normalised into a consistent ordering. Duplicate indices, noninteger entries, out-of-range values, missing boundaries and invalid dimensions cause errors. The implementation can therefore accept irregular internal spacing while refusing extrapolation outside acquired boundaries.

An acquired plane is copied from the normalised API input into the assembled output after optional smoothing. Tests and end-to-end checks assert exact equality, not approximate equality. This is stronger than an acquired-plane PSNR threshold and avoids obscuring small changes. It does not constrain the anatomical correctness of generated planes between acquired depths.''')
heading('5.4 CHECKPOINTS RESUME AND EMA',2)
prose('''The original best diffusion checkpoint is stored at epoch 48 with noise-validation objective 0.1177198. Required named states are loaded strictly. A missing or shape-incompatible checkpoint must terminate the operation. The original checkpoints are retained, and the review's candidate is saved in a separate directory with a manifest and SHA-256 identities.

Fresh initialisation from a checkpoint loads trained model weights but starts a fresh optimiser. Resume instead restores optimiser state, AMP scaler, exponential moving averages, training progress and random states. The step target is cumulative. The saved optimiser learning rate is retained on resume, so specifying a CLI rate is not a silent way to change it. Reproduction must keep the data, architecture, seed, stride set and batch size consistent.

The historical EMA averages only the U-Net. Newly saved checkpoints include an averaged context encoder as well. Raw and EMA comparisons therefore need the checkpoint's actual state definition. It would be incorrect to describe the historical EMA as a complete average of every trainable component. Both available modes are compared in the validation grid.''')
heading('5.5 COMPUTE AND MEMORY HANDLING',2)
para('The implementation uses mixed precision for supported CUDA model operations, while schedules and fidelity calculations remain in float32. Attention uses the supported scaled dot product attention path without changing checkpoint parameter shapes. Batches of eight missing planes balance memory and execution in the recorded environment. The image cache is bounded at 128 entries. No claim of universal performance across devices follows from this single laptop experiment.')
heading('5.6 BASELINE NAMING AND EXPORTS',2)
prose('''The former B-spline baseline was found to implement linear interpolation followed by Gaussian depth smoothing with standard deviation 0.7. It was renamed smoothed linear interpolation. The old method name and class name remain compatibility aliases, with documentation that they do not implement a B-spline. This correction changes how a comparison is described, not the historical numerical output.

Single-channel NIfTI output is ordered height, width and depth. Multiple-channel output adds channels as a fourth axis. HDF5 output preserves the reconstruction's depth, channels, height and width order. Tests checked one-channel and four-channel array export. The demonstrated trained model still uses a single channel.

When no affine is supplied, NIfTI qform and sform codes are unknown and spatial units are not asserted. Identity coordinates represent voxel indices only. JSON sidecars record the intensity convention, model settings and provenance. These are research exports; they are not verified patient-space images suitable for direct clinical navigation or measurement.''')
heading('5.7 VERIFICATION EVIDENCE',2)
table('Recorded verification results',['Check','Recorded outcome'],[
('Existing and focused tests','All 34 passed'),('Short training','Two steps completed'),('Checkpoint resume','Continued to cumulative step three'),('Repeated reconstruction','Volume 14, stride 8, seed 1450; maximum difference zero'),('Acquired planes','Exactly preserved in normalised input space'),('Index and checkpoint failures','Focused validation checks passed'),('Exports','One and four channels checked in NIfTI and HDF5')],[5.2,9.45])
para('These are the verification results recorded during the project review. They do not constitute new test execution during report writing, external certification or clinical validation. The appendices reproduce the critical source paths so that an examiner can inspect how these invariants are implemented. The available test evidence covers focused behaviours and end-to-end operation in the recorded environment.')

heading('CHAPTER 6 EXPERIMENTAL PROTOCOL')
heading('6.1 RESEARCH QUESTIONS',2)
prose('''The experiment asks whether a corrected conditional diffusion reference is competitive with interpolation at strides 8 and 16, whether bounded fine-tuning improves missing-slice fidelity, and which available inference settings perform best on the fixed validation cases. It also checks whether acquired data remain unchanged and whether the experiment can be reproduced from saved settings and model identities.

A corrected reference is necessary because fixing preprocessing can change metrics independently of learning. The experiment therefore compares the candidate with the same corrected evaluation procedure. Older results are retained for provenance but excluded from the final ranking and paired comparison.''')
heading('6.2 BUDGET AND TRAINING SETTINGS',2)
table('Allocated and measured protocol duration',['Stage','Allocation min','Measured min'],[
('Reference and validation',20,'10.67'),('Fine-tuning',70,'18.49'),('Selection and final comparison',30,'11.21'),('Total',120,'40.36')],[7.15,3.75,3.75])
para('The experiment remained inside the maximum two-hour budget. Fine-tuning stopped at the 3,000-step cap rather than using the entire 70-minute allocation. Stage values are rounded, so their rounded sum may differ from the rounded total. Smoke verification was separate and short. The result should not be described as a full 70-minute fine-tuning run.')
table('Completed fine-tuning settings',['Setting','Value'],[
('Initialisation','Existing best diffusion checkpoint'),('Updated components','Context encoder and U-Net'),('Frozen component','Existing best VAE'),('Training strides','4, 8 and 16; equal-probability draws'),('Learning rate','0.00002'),('Optimiser','AdamW; weight decay 0.0001'),('Batch size and steps','8; 3,000'),('Sampling','With replacement; missing targets only'),('Draw count','24,000 training draws, not unique slices'),('Seed','42'),('Gradient clipping','Norm 1'),('EMA decay','0.999'),('Augmentation','Disabled'),('Stopping reason','step_limit')],[5.1,9.55])
heading('6.3 VALIDATION AND MODEL SELECTION',2)
prose('''Five fixed validation volumes are evaluated at both final strides. For each model, raw and available EMA weights are combined with 25 or 50 DDIM steps and smoothing enabled or disabled. This creates eight settings per model and ten volume-stride pairs per setting. The same pair-specific generation seeds are used for the reference and candidate comparisons.

Settings are ranked by mean missing-slice PSNR across the five volumes and both strides. A candidate must improve PSNR strictly and must not decrease mean missing-slice SSIM relative to the best corrected reference setting. This rule was chosen before final test comparison. It avoids replacing the reference merely because one metric or one stride improves.

The rule does not guarantee improvement on each individual volume or each stride. Averaging the two strides assigns equal weight to their mean per-volume metrics, despite different missing-plane counts. A larger study could introduce lesion-specific metrics, noninferiority margins or confidence intervals, but such analyses are not part of the completed selection.''')
heading('6.4 METRIC DEFINITIONS',2)
equation(row(mi('MSE'),mo('='),frac(mn(1),mi('N')),mi('Σ'),sup(row(mo('('),sub('y','i'),mo('−'),sub('ŷ','i'),mo(')')),2)),'6.1')
equation(row(mi('PSNR'),mo('='),mn(10),mo('·'),sub('log','10'),mo('('),frac(mn(4),mi('MSE')),mo(')')),'6.2')
prose('''For missing-slice PSNR, squared error is pooled over pixels in all missing planes for a volume, using normalised data range two. The factor four is the square of that range. If MSE is at or below one ten-billionth, the implementation returns a sentinel of 100 dB. This avoids numerical infinity and must be recognised when interpreting identical arrays.

SSIM is computed per two dimensional slice and channel, then averaged. The implementation uses an eleven by eleven Gaussian window with standard deviation 1.5 and zero padding. Its constants are the squares of 0.01 times two and 0.03 times two. Full-volume and missing-slice metrics include background. No lesion mask or clinical-region weighting is applied.

Full-volume scores can be inflated by exact acquired slices, so selection uses missing-slice metrics. PSNR measures squared intensity error; SSIM measures a different local structural relationship. Their rankings can disagree. This report presents both and avoids treating either as a proof of diagnostic correctness.''')
heading('6.5 REPRODUCIBILITY AND TIMING',2)
equation(row(sub('seed','reconstruction'),mo('='),mn(42),mo('+'),mn(100),mo('·'),sub('volume','id'),mo('+'),mi('stride')),'6.3')
prose('''The recorded reconstruction seed combines the global seed, volume ID and stride. The repeated reconstruction check used volume 14 and stride 8, giving seed 1450. Saved manifests record device, precision, settings, checkpoint hashes and elapsed stage times. A repeated local run with the final reviewed source yielded maximum absolute difference zero for the checked volume.

Diffusion row timings cover generation with CUDA synchronisation and exclude smoothing and metric calculation. Baseline timings include interpolation and metrics. They are not equal-work timing measurements. Stage manifests provide broader elapsed times, but neither type of timing is a scanner acquisition measurement. Smoothing settings reuse the same generated volume; their reported timings must not be added as separate generation calls.

Reproducibility is demonstrated for the recorded environment and checked setting. Different devices, library kernels or batching may affect numerical output. There was one fine-tuning seed and one candidate, so the experiment does not estimate variation across independently trained models.''')
heading('6.6 EVIDENCE FILES',2)
para('The evidence directory contains experiment_manifest.json, selection.json, per-volume validation and test CSVs, configuration snapshots, source checksums and paired comparisons. The fine-tuning directory contains manifest.json, training.csv, completion.json and a resumable latest.pt. Original checkpoints and historical result files were not overwritten. Appendix 1 gives commands and identities for inspecting and reproducing the work.')

heading('CHAPTER 7 RESULTS AND DISCUSSION')
heading('7.1 VALIDATION GRID',2)
rows=[]
for filename,label in [('validation_reference.csv','Reference'),('validation_candidate.csv','Fine-tuned')]:
    data=list(csv.DictReader((ROOT/'results/review'/filename).open()))
    groups={}
    for r in data:groups.setdefault((r['weights'],r['steps'],r['smoothing']),[]).append(r)
    for key,rs in sorted(groups.items()):rows.append((label,*key,f"{sum(float(x['psnr_missing']) for x in rs)/len(rs):.4f}",f"{sum(float(x['ssim_missing']) for x in rs)/len(rs):.5f}"))
table('Validation means over five volumes and two strides',['Model','Weights','Steps','Smooth','PSNR dB','SSIM'],rows,[3.1,1.9,1.8,1.8,3.05,3])
prose('''The corrected reference's best setting is raw weights, 25 steps and smoothing enabled, with mean missing-slice PSNR 26.1862 dB and SSIM 0.89114. The best candidate by PSNR uses EMA weights, 25 steps and smoothing, with PSNR 26.1643 dB and SSIM 0.88992. The candidate is lower on both means and therefore fails the eligibility rule.

Smoothing improved both validation metrics for every evaluated setting in this grid. Fifty steps were not better than 25 for the compared models and weight modes. These observations apply to this checkpoint, dataset subset and seed schedule. They do not establish a general rule about diffusion inference length or smoothing.

Raw and EMA reference scores are close, while the candidate's EMA improves slightly over its raw state at 25 smoothed steps. This does not rescue candidate eligibility. The selected checkpoint is the corrected original reference; no learned improvement was accepted during this bounded attempt.''')
heading('7.2 FINAL BENCHMARK COMPARISON',2)
data=list(csv.DictReader((ROOT/'results/review/test_reference.csv').open()))
import statistics as S
summary=[]
for stride in (8,16):
    for method in ['conditional_diffusion','linear','smoothed_linear','cubic','nearest']:
        rs=[r for r in data if int(r['stride'])==stride and r['method']==method]
        if not rs:continue
        ps=[float(r['psnr_missing']) for r in rs];ss=[float(r['ssim_missing']) for r in rs]
        summary.append((stride,('Diffusion' if method=='conditional_diffusion' else method.replace('_',' ')),f'{S.mean(ps):.3f} ± {S.stdev(ps):.3f}',f'{S.mean(ss):.4f} ± {S.stdev(ss):.4f}'))
table('Missing-slice fidelity over ten benchmark volumes',['Stride','Method','PSNR dB mean ± SD','SSIM mean ± SD'],summary,[1.9,3.25,4.75,4.75])
figure(Path('results/review/measured_results.png'),'Missing-slice metrics with between-volume standard deviation')
prose('''Diffusion achieved mean PSNR 28.037 dB at stride 8 and 25.071 dB at stride 16. Linear interpolation achieved 27.628 dB and 24.550 dB respectively. The corresponding diffusion advantages were approximately 0.409 dB and 0.521 dB. These are dataset-level fidelity differences under the corrected preprocessing, not clinical effect sizes.

At stride 8, linear interpolation had higher mean SSIM: 0.9226 compared with diffusion's 0.9054. At stride 16, diffusion had higher mean SSIM: 0.8818 compared with linear interpolation's 0.8752. The result is therefore mixed across metrics and gap sizes. A statement that diffusion is uniformly superior would contradict the benchmark.

The selected and corrected reference results are identical in quality because the same model and settings were retained. Their repeated timing measurements differ. Mean paired changes in both PSNR and SSIM are zero, with zero standard deviation, across twenty volume-stride pairs. This is the correct conclusion of the selection experiment, rather than an improvement claim for fine-tuning.''')
heading('7.3 PER VOLUME RESULTS',2)
for stride in (8,16):
    rs=sorted([r for r in data if int(r['stride'])==stride and r['method']=='conditional_diffusion'],key=lambda x:int(x['vol_id']))
    table(f'Corrected diffusion per-volume results at stride {stride}',['Volume','PSNR dB','SSIM','Generation s'],[(r['vol_id'],f"{float(r['psnr_missing']):.3f}",f"{float(r['ssim_missing']):.4f}",f"{float(r['seconds']):.2f}") for r in rs],[2.4,4.1,4.05,4.1])
para('The reported standard deviations describe variability among the ten local benchmark volumes. They are not confidence intervals and do not measure variability across repeated training seeds. Per-volume values are provided so that a reader can inspect heterogeneity rather than relying only on a pooled headline number. No statistical significance claim is made.')
heading('7.4 QUALITATIVE EXAMPLE',2)
figure(Path('results/review/example_reconstruction_montage.png'),'Saved reconstruction montage for volume 14 at stride 8',width=14.4)
para('The saved montage illustrates axial output planes at multiple depths for a representative reconstruction. It is an output illustration, not a controlled side-by-side diagnostic comparison. Some displayed depths coincide with acquired planes while others are generated. A visually coherent shape can still contain incorrect tissue detail, and the montage is not evidence of lesion preservation or clinical suitability.')
heading('7.5 COMPUTE OBSERVATIONS',2)
table('Recorded mean row timings with different measured scopes',['Stride','Reference diffusion s','Selected diffusion s','Linear s'],[(8,'7.65','7.93','0.15'),(16,'7.29','8.40','0.15')],[1.7,4.8,4.8,3.35])
para('These timings are included for transparency. Diffusion measures generation; linear measures interpolation plus metrics. Consequently the table does not support a fair numerical speedup ratio. Model loading, preprocessing, export and clinical workflow overhead also require separate accounting. The actual complete bounded protocol took 40.36 minutes, and fine-tuning accounted for 18.49 minutes of it.')
heading('7.6 WHY THE CANDIDATE WAS NOT RETAINED',2)
prose('''The fine-tuning result shows that a plausible training recipe does not guarantee better reconstruction. Noise-prediction training loss and final decoded image metrics measure different quantities. A frozen VAE can constrain fidelity, and the corrected acquired-only intensity distribution may differ from the VAE's historical training distribution. These are possible explanations, not experimentally isolated causes.

The attempt used one learning rate, one seed, one starting checkpoint and a 3,000-step cap. No learning-rate sweep, VAE adaptation or expanded training study was performed. It would be incorrect to attribute the failed selection to any single factor without an ablation. Retaining the reference protects the reported result from an unsupported improvement claim while preserving the candidate for further analysis.''')
heading('7.7 LIMITATIONS AND THREATS TO VALIDITY',2)
prose('''First, the benchmark is a small internal subset of ten local volumes and two acquisition patterns. It does not establish generalisation to another institution, scanner, contrast or disease distribution. The original patient mapping and conversion provenance are incompletely verified. The split is disjoint by local volume ID, which is necessary but weaker than an independently verified clinical grouping.

Second, metrics include background and do not isolate lesions or clinically important boundaries. Mean slice SSIM and pooled missing-plane PSNR can obscure regional failure. The sampled mask semantics are not verified, so adding lesion or Dice evaluation requires a provenance step before metric implementation.

Third, the model generates planes using two dimensional processing and context. Gaussian smoothing and exact acquired restoration can improve continuity of intensities, but they do not enforce a physical three dimensional anatomical model. A missing small lesion cannot be guaranteed to reappear from bounding observations.

Fourth, physical geometry is absent from the converted data. Exports carry unknown orientation and spacing rather than verified patient coordinates. Clinical integration would require reliable original geometry, units, metadata handling and a validated acquisition protocol.

Finally, this is a single bounded learning attempt. The selected reference provides the best measured setting under the specified validation rule, but the experiment does not exhaust possible architectures or training recipes. Historical scores with different preprocessing and claims from proposal slides cannot be combined with the corrected benchmark. No clinical benefit, scan-time reduction or patentability conclusion is inferred.''')

heading('CHAPTER 8 CONCLUSION AND FUTURE WORK')
heading('8.1 CONCLUSION',2)
prose('''The project delivers a working research pipeline for reconstructing missing brain MRI slices using conditional latent diffusion. The implementation combines shared context encoding, explicit coordinate conditioning, a frozen custom VAE, a latent denoising U-Net and seeded DDIM inference. It assembles a dense output volume while preserving the supplied acquired planes exactly in normalised space.

The review corrected preprocessing dependencies, missing-target selection, irregular-index handling, trained-checkpoint requirements, resume and EMA state, baseline terminology and export semantics. These corrections improve the reliability and interpretation of experiments even though they do not by themselves establish a learned performance gain. The recorded verification includes 34 passing tests, a short training and resume check, export checks and an exactly repeated reconstruction.

The bounded performance attempt completed inside two GPU hours but stopped fine-tuning at 3,000 steps. The candidate failed the validation rule. The corrected original model was retained with raw weights, 25 DDIM steps and smoothing. On the ten saved benchmark volumes, diffusion had higher mean missing-slice PSNR than linear interpolation at both strides, while SSIM favoured linear interpolation at stride 8 and diffusion at stride 16. The evidence supports a reproducible prototype with mixed fidelity benefits and substantial remaining validation needs.''')
heading('8.2 PRIORITISED FUTURE STUDIES',2)
prose('''The first priority is provenance: verify channel identities, original patient mapping, mask semantics and physical geometry against the source data. This enables reliable lesion-aware metrics and spatially meaningful exports. A separate held-out cohort should then assess cross-site and cross-scanner generalisation, with acquisition patterns defined by actual available slices.

A training study should compare additional learning rates and multiple seeds, with predeclared validation criteria and matched compute. VAE adaptation should be investigated separately from denoiser adaptation because it changes the latent representation and invalidates direct checkpoint compatibility. Each ablation should record both improvement and degradation, including acquired-plane fidelity and regional errors.

Three dimensional or neighbouring-plane consistency objectives could be studied to reduce inconsistent generated structures. An uncertainty study could evaluate variation across seeds and identify whether dispersion correlates with anatomical error. Neither smoothing nor a single stochastic sample should be described as a validated uncertainty estimate.

A clinical study would require domain specialists, prospective or independently held-out acquisition data, suitable ethical and institutional approvals, and explicit diagnostic endpoints. Scanner acquisition-time savings would need a real acquisition protocol and measurement. CT, other organs, joint multicontrast models and downstream segmentation remain separate research proposals until such evidence is produced.''')
heading('8.3 REPRODUCIBLE HANDOVER',2)
para('The source tree, trained artifacts, corrected experiment CSVs and manifests form the technical handover. The README explains portable setup and compatibility entry points. The review findings and this report distinguish implemented features from proposed uses. The appendices below provide exact module identities, experiment controls and selected source listings so that future work can proceed from the measured reference without overwriting historical evidence.')

heading('REFERENCES')
refs=[
'Bubeck, N., Shit, S., Chen, C., Zhao, C., Guo, P., Yang, D., Zitzlsberger, G., Xu, D., Kainz, B., Rueckert, D., & Pan, J. (2025). Latent interpolation learning using diffusion models for cardiac volume reconstruction (Version 4, revised July 24, 2026) [Preprint]. arXiv. https://doi.org/10.48550/arXiv.2508.13826',
'Center for Biomedical Image Computing and Analytics. (2020). BraTS 2020 data. University of Pennsylvania. https://www.med.upenn.edu/cbica/brats2020/data.html',
'Ho, J., Jain, A., & Abbeel, P. (2020). Denoising diffusion probabilistic models. Advances in Neural Information Processing Systems, 33, 6840–6851. https://arxiv.org/abs/2006.11239',
'Kingma, D. P., & Welling, M. (2014). Auto-encoding variational Bayes. International Conference on Learning Representations. https://arxiv.org/abs/1312.6114',
'Nichol, A. Q., & Dhariwal, P. (2021). Improved denoising diffusion probabilistic models. Proceedings of Machine Learning Research, 139, 8162–8171. https://proceedings.mlr.press/v139/nichol21a.html',
'Rombach, R., Blattmann, A., Lorenz, D., Esser, P., & Ommer, B. (2022). High-resolution image synthesis with latent diffusion models. Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition, 10684–10695. https://arxiv.org/abs/2112.10752',
'Ronneberger, O., Fischer, P., & Brox, T. (2015). U-Net: Convolutional networks for biomedical image segmentation. In N. Navab, J. Hornegger, W. M. Wells, & A. F. Frangi (Eds.), Medical image computing and computer-assisted intervention – MICCAI 2015 (pp. 234–241). Springer. https://doi.org/10.1007/978-3-319-24574-4_28',
'Song, J., Meng, C., & Ermon, S. (2021). Denoising diffusion implicit models. International Conference on Learning Representations. https://arxiv.org/abs/2010.02502'
]
ref_titles=['Latent interpolation learning using diffusion models for cardiac volume reconstruction','BraTS 2020 data','Advances in Neural Information Processing Systems','Auto-encoding variational Bayes','Proceedings of Machine Learning Research','Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition','Medical image computing and computer-assisted intervention – MICCAI 2015','Denoising diffusion implicit models']
for ref,italic_title in zip(refs,ref_titles):
    p=para('',WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7)
    a,b=ref.split(italic_title,1);p.add_run(a);p.add_run(italic_title).italic=True;p.add_run(b)

heading('APPENDIX 1 REPRODUCTION AND EVIDENCE REGISTER')
heading('A1.1 LOCAL ENVIRONMENT AND COMMANDS',2)
para('Run commands from the project directory using the supplied Python environment. The commands below are documented interfaces; reproduction should use separate output directories to preserve the completed records. The complete review protocol has fixed default output paths and should only be rerun after preserving them or changing the script paths.')
commands=[
"$projectPython = 'D:\\Kaggle competetion\\torch_env\\Scripts\\python.exe'",
'& $projectPython scripts/reconstruct_volume.py --vol_id 14 --stride 8 --seed 42 --steps 25 --weights raw --output-dir results/review/new_reconstruction',
'& $projectPython scripts/reconstruct_volume.py --vol_id 14 --stride 16 --format h5 --no-smoothing --output-dir results/review/new_hdf5',
'& $projectPython scripts/evaluate_all.py --vols 10 --strides 8 16 --output-dir results/review/new_benchmark',
'& $projectPython scripts/finetune_diffusion.py --init-checkpoint experiments/checkpoints/diffusion/best_model.pt --strides 4 8 16 --lr 2e-5 --seed 42 --time-limit-minutes 70 --max-steps 3000 --batch-size 8 --output-dir experiments/review/new_run',
'& $projectPython scripts/finetune_diffusion.py --resume experiments/review/new_run/latest.pt --strides 4 8 16 --seed 42 --max-steps 6000 --time-limit-minutes 30 --output-dir experiments/review/new_run',
'& $projectPython -m pytest tests -q',
'& $projectPython scripts/review_experiment.py'
]
for i,c in enumerate(commands,1):
    para(f'Command {i}',bold=True);para(c,WD_ALIGN_PARAGRAPH.LEFT)
heading('A1.2 TRAINED ARTIFACT IDENTITIES',2)
for name,h in [('Diffusion best_model.pt','9644972de899da7200f179a0f940899730c17a5f65edeeaddd8083aa779cac52'),('VAE best_model.pt','0b937d72ccecd68460b1ac0660e7296d711f27c8a50a83f043eedfd8abb030ff')]:
    para(name,bold=True);para('SHA-256: '+h,WD_ALIGN_PARAGRAPH.LEFT)
para('Paths: experiments/checkpoints/diffusion/best_model.pt and experiments/checkpoints/vae/best_model.pt. The candidate is experiments/review/finetune/latest.pt. The selected checkpoint is the original diffusion checkpoint after evaluation corrections, not the saved candidate.')
table('Evidence register',['File or directory','Evidence role'],[
('results/review/dataset_audit.json','Collection counts, shapes and metadata audit'),('results/review/REVIEW_FINDINGS.md','Corrections and remaining limitations'),('results/review/experiment_manifest.json','Budget, cases and seed'),('results/review/selection.json','Selected checkpoint and eligibility outcome'),('results/review/test_reference.csv','Per-volume reference and baseline metrics'),('results/review/test_selected.csv','Selected-model metrics'),('results/review/validation_reference.csv','Reference setting grid'),('results/review/validation_candidate.csv','Candidate setting grid'),('results/review/paired_comparison.csv','Paired test metric differences'),('results/review/reproducibility_check.json','Repeated reconstruction and acquired equality'),('experiments/review/finetune/','Run settings, losses, completion and checkpoint')],[7.5,7.15])
heading('A1.3 CONFIGURATION INTERPRETATION',2)
para('The default and diffusion configuration files below document the architecture and supported settings. Some historical training fields remain in YAML for compatibility and do not govern the bounded fine-tuning CLI. The completed-run manifest is the authority for actual learning rate, step count, batch size, stride sampling and stopping reason. Code comments are reproduced verbatim in Appendix 3; broad historical MRI/CT or DDPM labels in comments are not claims that those uses were validated.')
for path in ['configs/default_config.yaml','configs/diffusion_config.yaml']:
    heading(path.replace('/',' ').replace('_',' ').upper(),3)
    for line in (ROOT/path).read_text(encoding='utf-8').splitlines():
        p=para(line or ' ',WD_ALIGN_PARAGRAPH.LEFT,after=0)
        p.paragraph_format.widow_control=False

heading('APPENDIX 2 PAIRED BASELINE RECORDS')
para('These tables reproduce the corrected reference and linear-interpolation values for every test volume and stride. They use missing planes only. The selected diffusion values equal the corrected reference; a second identical column is unnecessary. Values are rounded for reading; the CSVs retain full recorded precision.')
for stride in (8,16):
    rs=[]
    for vol in [14,21,22,35,49,50,51,53,55,59]:
        dr=next(r for r in data if int(r['stride'])==stride and int(r['vol_id'])==vol and r['method']=='conditional_diffusion')
        lr=next(r for r in data if int(r['stride'])==stride and int(r['vol_id'])==vol and r['method']=='linear')
        rs.append((vol,f"{float(dr['psnr_missing']):.3f}",f"{float(lr['psnr_missing']):.3f}",f"{float(dr['ssim_missing']):.4f}",f"{float(lr['ssim_missing']):.4f}"))
    table(f'Paired diffusion and linear metrics at stride {stride}',['Volume','Diffusion PSNR','Linear PSNR','Diffusion SSIM','Linear SSIM'],rs,[2.05,3.15,3.15,3.15,3.15])

heading('APPENDIX 3 CRITICAL SOURCE LISTINGS')
para('The following modules are reproduced from the reviewed repository for inspection of the executed reconstruction, scheduling, data validation, metric and export paths. Line numbers identify the corresponding source lines; blank source lines are omitted and automatic Word wrapping is presentational only. Listings are snapshots, not a substitute for executable files. Source hashes allow comparison with subsequent changes. Comments retain their original wording and must be read alongside the demonstrated-scope limits in the main report.')
source_paths=['src/data/preprocessing.py','src/data/sparsifier.py','src/data/transforms.py','src/models/factory.py','src/models/pipeline.py','src/models/scheduler.py','src/evaluation/metrics.py','src/utils/volume_io.py','tests/test_review.py']
snapshots=[]
for k,path in enumerate(source_paths,1):
    hp=heading(f'A3.{k} '+path.replace('/',' ').replace('_',' ').upper(),2)
    hp.paragraph_format.space_before=Pt(6);hp.paragraph_format.space_after=Pt(4)
    blob=(ROOT/path).read_bytes();h=hashlib.sha256(blob).hexdigest()
    para('Source: '+path,WD_ALIGN_PARAGRAPH.LEFT,after=0)
    para('SHA-256: '+h,WD_ALIGN_PARAGRAPH.LEFT,after=0)
    lines=blob.decode('utf-8-sig').splitlines()
    snapshots.append({'file':path,'sha256':h,'lines':len(lines)})
    for i,line in enumerate(lines,1):
        if not line.strip():continue
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=1.5;p.paragraph_format.widow_control=False
        p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True
        # Preserve indentation without allowing Word to collapse leading spaces.
        line=line.expandtabs(4)
        indent=len(line)-len(line.lstrip(' '))
        p.add_run(f'{i:03d}  '+('\u00a0'*indent)+line[indent:])

def insert_list(anchor,records):
    last=anchor._p
    for title,name in records:
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        p.add_run(title+'\t');field(p,' PAGEREF '+name+' \\h ','0')
        last.addnext(p._p);last=p._p
insert_list(table_list_anchor,tables)
insert_list(figure_list_anchor,figures)
doc.core_properties.title=TITLE.title()
doc.core_properties.author='Gudikandula Sai Charan Chary'
doc.core_properties.subject='Individual project report for 23BAI1536'
doc.core_properties.keywords='BraTS, FLAIR, latent diffusion, sparse reconstruction'
doc.save(FINAL)
audit={'output':str(FINAL),'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),'tables':len(tables),'figures':len(figures),'native_math':len(doc._element.xpath('.//m:oMath')),'source_snapshots':snapshots,'external_guide_certificate':False,'submission':'October 2026','render_required':True}
(OUT/'build_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print(json.dumps(audit,indent=2))
