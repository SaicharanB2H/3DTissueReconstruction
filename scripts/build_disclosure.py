"""Create the disclosure and drawings from verified local experiment artifacts."""
import sys
import json
import re
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.review_tools'))
sys.path.insert(0, str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = ROOT / 'results/review'
DRAWINGS = OUT / 'drawings'
DRAWINGS.mkdir(parents=True, exist_ok=True)

def drawing(name, nodes, arrows, labeled=True, size=(12,7)):
    fig, ax = plt.subplots(figsize=size)
    ax.set_xlim(0,12)
    ax.set_ylim(0,8)
    ax.axis('off')
    for x,y,w,h,label in nodes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06',facecolor='white',edgecolor='black',linewidth=1.5))
        if labeled:
            ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=18 if name=='architecture' else 16)
    for a,b in arrows:
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=16,color='black',linewidth=1.3))
    fig.tight_layout(pad=.5)
    base=DRAWINGS/(name+('_labeled' if labeled else '_unlabeled'))
    fig.savefig(str(base)+'.png',dpi=180)
    fig.savefig(str(base)+'.svg')
    plt.close(fig)

def diagrams():
    architecture=[
        (.3,6.4,2.4,1.1,'101 Sparse slices\n+ depth indices'),
        (3.2,6.4,2.4,1.1,'102 Acquired-only\nnormalization'),
        (6.1,6.4,2.5,1.1,'103 Shared context\nencoder'),
        (9.1,6.4,2.5,1.1,'104 Fourier\ncoordinates'),
        (6.1,3.9,2.5,1.1,'105 Tokens\n+ spatial maps'),
        (.3,3.9,2.4,1.1,'106 Seeded latent\nGaussian noise'),
        (3.2,3.9,2.4,1.1,'107 U-Net\n+ DDIM updates'),
        (3.2,1.3,2.4,1.1,'108 VAE decoder\n(frozen)'),
        (6.1,1.3,2.5,1.1,'109 Assembly\n+ preservation'),
        (9.1,1.3,2.5,1.1,'110 Volume\n+ export')]
    arrows=[((2.7,6.95),(3.2,6.95)),((5.6,6.95),(6.1,6.95)),((7.35,6.4),(7.35,5.0)),
        ((10.35,6.4),(8.6,4.45)),((6.1,4.45),(5.6,4.45)),((2.7,4.45),(3.2,4.45)),
        ((4.4,3.9),(4.4,2.4)),((5.6,1.85),(6.1,1.85)),((8.6,1.85),(9.1,1.85)),
        ((4.4,6.4),(2.95,5.7)),((2.95,5.7),(2.95,2.9)),((2.95,2.9),(7.35,2.9)),((7.35,2.9),(7.35,2.4))]
    # Calibrated acquired values bypass generation and are restored by assembly.
    flow=[(3,6.8,6,0.8,'201 Validate volume, channels\nand acquired indices'),
          (3,5.55,6,.8,'202 Normalize intensities\nfrom acquired slices only'),
          (3,4.3,6,.8,'203 Enumerate missing depths\nand bounding contexts'),
          (3,3.05,6,.8,'204 Encode context; run\nconditional latent DDIM'),
          (3,1.8,6,.8,'205 Decode; optionally smooth;\nrestore acquired slices'),
          (3,.55,6,.8,'206 Export volume + provenance;\nevaluate when ground truth exists')]
    flow_arrows=[((6,y),(6,y-.45)) for y in [6.8,5.55,4.3,3.05,1.8]]
    for labeled in (True,False):
        drawing('architecture',architecture,arrows,labeled)
        drawing('operation',flow,flow_arrows,labeled,size=(9,10))

SOURCES=[
 ('R1','Rombach et al. High-Resolution Image Synthesis with Latent Diffusion Models (2021/2022)','https://arxiv.org/abs/2112.10752'),
 ('R2','Bubeck et al. Latent Interpolation Learning Using Diffusion Models for Cardiac Volume Reconstruction (first posted 19 Aug 2025)','https://arxiv.org/abs/2508.13826'),
 ('R3','Du et al. Improving 2D Diffusion Models for 3D Medical Imaging with Inter-Slice Consistent Stochasticity (2026)','https://arxiv.org/abs/2602.04162'),
 ('R4','US11410271B2, Reconstructing missing slices in slice-based scans of 3D objects (9 Aug 2022); claims 1, 11, 14 reviewed','https://patents.google.com/patent/US11410271B2/en'),
 ('R5','WO2025264018A1, multi-slice-trained 2D diffusion for 3D volume conversion: bibliographic lead; full text retrieval unavailable','https://patents.google.com/patent/WO2025264018A1/en'),
 ('R6','Official BraTS 2020 imaging and annotation description','https://www.med.upenn.edu/cbica/brats2020/data.html'),
 ('R7','Siemens Healthcare GmbH official contact address','https://www.siemens-healthineers.com/en-th/services/value-partnerships/contact-us'),
 ('R8','Siemens Deep Resolve manufacturer description','https://www.siemens-healthineers.com/en-au/magnetic-resonance-imaging/technologies-and-innovations/deep-resolve'),
 ('R9','Koninklijke Philips N.V. 2025 SEC filing: principal address','https://www.sec.gov/Archives/edgar/data/313216/000162828026009470/R1.htm'),
 ('R10','Philips 2025 annual results','https://www.results.philips.com/ar25'),
 ('R11','Philips Diagnosis & Treatment investor presentation (2026)','https://www.philips.com/c-dam/corporate/about-philips/investors/calendar/2026/3_-_Diagnosis_and_Treatment.pdf')]

def results_text():
    if not (OUT/'completion.json').exists():
        return 'Corrected experiments are in progress. Final measured results will be inserted after completion; no improvement is asserted yet.'
    completion=json.loads((OUT/'completion.json').read_text())
    selection=json.loads((OUT/'selection.json').read_text())
    ref=pd.read_csv(OUT/'test_reference.csv')
    selected=pd.read_csv(OUT/'test_selected.csv')
    parts=['Corrected local evaluation on ten held-out volumes (14, 21, 22, 35, 49, 50, 51, 53, 55, 59), using acquired-only normalization and fixed seeds.']
    for stride in [8,16]:
        s=selected[selected.stride==stride]
        r=ref[(ref.stride==stride)&(ref.method=='conditional_diffusion')]
        linear=ref[(ref.stride==stride)&(ref.method=='linear')]
        parts.append(f'Stride {stride}: selected missing-slice PSNR {s.psnr_missing.mean():.3f} +/- {s.psnr_missing.std():.3f} dB, SSIM {s.ssim_missing.mean():.4f} +/- {s.ssim_missing.std():.4f}; corrected raw 25-step reference PSNR {r.psnr_missing.mean():.3f} dB, SSIM {r.ssim_missing.mean():.4f}; linear PSNR {linear.psnr_missing.mean():.3f} dB, SSIM {linear.ssim_missing.mean():.4f}.')
    parts.append(f'Validation selected {"fine-tuned candidate" if selection["use_candidate"] else "original reference"}, {selection["weights"]} weights, {selection["steps"]} DDIM steps, smoothing {selection["smoothing"]}. Total experiment elapsed time {completion["seconds"]/60:.1f} minutes. Metrics measure image fidelity including background, not diagnosis or lesion segmentation. No independent third-party clinical validation is available.')
    return '\n'.join(parts)

def answers():
    training_result = json.loads((ROOT/'experiments/review/finetune/completion.json').read_text()) if (ROOT/'experiments/review/finetune/completion.json').exists() else None
    training_note = (f" Fine-tuning completed {training_result['step']} training iterations ({training_result['step']*8} sampled slice examples with replacement) in {training_result['seconds']/60:.2f} minutes, stopping at the configured step cap within the 70-minute limit. The candidate did not satisfy the validation improvement criterion, so the corrected original checkpoint remains selected." if training_result and (OUT/'selection.json').exists() and not json.loads((OUT/'selection.json').read_text())['use_candidate'] else '')
    technical='A computer-implemented research system reconstructs missing axial brain MRI slices from sparse acquired slices. A shared residual convolutional encoder processes the bounding slices. Three-coordinate Fourier embeddings encode normalized depth, relative position and slice gap. Spatial context and cross-attention tokens condition a latent denoising U-Net; DDIM sampling and a pretrained VAE decoder generate missing slices. Assembly optionally smooths along depth and restores acquired slices exactly.'
    working='Input: a registered depth-first stack (D,C,H,W) and unique acquired depth indices including both boundaries. Validate channels, dimensions, finite values and checkpoints. Derive 0.5th/99.5th percentile foreground intensity bounds only from acquired slices; apply the same clipping/scaling to context and targets. For each missing depth locate bounding acquired slices and calculate z/(D-1), alpha and gap/16. Encode context, seed a Gaussian latent, predict noise through 25 or 50 DDIM updates, decode and assign the generated slice. Optionally apply Gaussian smoothing (sigma 0.5 along depth) and restore original acquired values. Export a research NIfTI/HDF5 volume with normalization, seed and checkpoint provenance; evaluate against withheld targets only when available.'
    prior='Established latent diffusion and cross-attention: R1. Sparse cardiac MRI latent reconstruction: CaLID, R2, a close conceptual overlap. Inter-slice noise consistency: ISCS, R3. Missing-slice binary interpolation: US11410271B2, R4. Multi-slice 2D diffusion/3D conversion patent lead: R5 (full text unavailable). See attached source register and PRIOR_ART.md. No head-to-head evaluation against these learned systems was performed.'
    equations='Coordinates: alpha=(z-z_prev)/(z_next-z_prev), depth=z/(D-1), gap=(z_next-z_prev)/16. Normalization: I_norm=2*(clip(I,vmin,vmax)-vmin)/(vmax-vmin+epsilon)-1, with background mapped to -1 and epsilon=1e-7. Clean latent: z0=0.18215*(mu+sigma*epsilon). Forward diffusion: z_t=sqrt(alpha_bar_t)*z0+sqrt(1-alpha_bar_t)*epsilon. Noise objective: E[||epsilon-epsilon_theta(z_t,t,context)||^2]. Deterministic DDIM: predicted z0=(z_t-sqrt(1-alpha_bar_t)*epsilon_theta)/sqrt(alpha_bar_t); z_prev=sqrt(alpha_bar_prev)*predicted z0+sqrt(1-alpha_bar_prev)*epsilon_theta. VAE objective: L1+0.5*MSE+1e-5*KL+0.1*spatial-gradient L1. Reconstruction MSE=mean((prediction-target)^2); MAE=mean(abs(prediction-target)); PSNR=10*log10(4/MSE) for [-1,1], capped at 100 dB for near-perfect matches. SSIM uses an 11x11 Gaussian window (sigma 1.5), constants (0.01*2)^2 and (0.03*2)^2, averaged over 2D slices/channels.'
    return {
    1:'Depth-Conditioned Latent Diffusion Reconstruction of Dense Brain MRI Volumes from Sparse Slices',
    2:'Problem: sparse axial sampling leaves unobserved planes and can obscure structures during volume resampling. Solution: '+technical+' The demonstrated scope is retrospective FLAIR MRI reconstruction; scanner acquisition acceleration and diagnostic efficacy are proposed applications requiring separate validation.',
    3:'Sparse-to-dense MRI; missing-slice synthesis; through-plane super-resolution; conditional latent diffusion; neighboring-slice conditioning; Fourier depth embedding; hybrid spatial/cross-attention conditioning; DDIM; acquired-slice preservation; volume interpolation; intensity calibration.',
    4:'Present methods include nearest/linear/cubic interpolation, learned slice synthesis, latent diffusion (R1/R2), and approaches addressing inter-slice stochasticity (R3). Commercial AI reconstruction such as Deep Resolve (R8) also exists. Classical interpolation is inexpensive but limited by sampled anatomy; learned priors can generate plausible unsupported detail. The implemented system has no k-space/sinogram measurement model, explicit 3D consistency loss, calibrated uncertainty or clinical validation.',
    5:prior,
    6:'A practical implementation gap is addressed: reproducible MRI slice reconstruction on a local 8 GB GPU with acquired-only calibration and explicit data preservation. A major gap in prior art has not been established; CaLID substantially overlaps sparse MRI latent interpolation. The specific combined conditioning design needs further full-text comparison and an inventor-confirmed disclosure timeline.',
    7:'Inputs and metadata validator; HDF5 slice reader; acquired-only intensity normalizer; depth/context extractor; weight-tied residual convolutional encoder (base channels 32); three-coordinate Fourier embedding; 17x256 context tokens and 8x30x30 spatial maps; pretrained MedicalAutoencoderKL (4x30x30 latent); conditional U-Net (base channels 64, multipliers 1/2/4); 1000-step noise scheduler; DDIM sampler; batch assembly and optional depth smoothing; NIfTI/HDF5 exporter; fidelity metrics and visualization modules. Software runs in Python/PyTorch on CPU or CUDA, with GPU used for the measured experiments.',
    8:working,
    9:'Candidate distinguishing combination: acquired-only intensity calibration plus weight-tied bounding-slice encoding, explicit global-depth/relative-offset/gap embeddings, hybrid spatial/token conditioning and exact acquired-slice restoration. These elements are implemented. Individual latent diffusion, cross-attention and slice interpolation concepts are already known (R1-R4). Novelty and inventive step remain unestablished; broad superiority or a first-in-field claim would exceed the evidence.',
    10:'Demonstrated operational features include bounded local GPU execution, strict checkpoint validation, reproducible seeds, missing-only evaluation and exact preservation of acquired planes. '+results_text()+' No comparative clinical, runtime, memory or cost advantage over CaLID, ISCS or commercial scanner software has been established.',
    11:'Proposed variants requiring implementation and validation: multiple MRI contrasts, CT with verified HU handling, other organs, physical-coordinate embeddings, larger context windows, correlated noise trajectories, learned inter-slice consistency, uncertainty ensembles and scanner-data consistency. Four-channel export is tested; existing trained weights are single-channel FLAIR. A CPU mode exists but CPU inference performance is not benchmarked.',
    12:'Four software schematics are supplied and embedded below: labeled/unlabeled architecture and labeled/unlabeled operation flow. Essential numbered components are explained in section 13. No physical-device CAD prototype is claimed. Historical Blender meshes are visualization artifacts and are not proof of reconstructed lesion accuracy.',
    13:'Figure 1: input acquired slices (101) are calibrated (102) and encoded (103); coordinates (104) contribute conditioning (105), which guides noise (106) through U-Net/DDIM (107). VAE decoding (108) produces missing planes; assembly (109) restores acquired planes and export (110) saves the volume. The bypass arrow carries acquired values to assembly. Figure 2 shows the same architecture without labels. Figure 3 depicts validation, acquired-only normalization, missing-context extraction, generation, preservation and export (201-206). Figure 4 repeats the flow without annotations.',
    14:technical+'\n'+working+'\nWorking example: D=155, C=1, H=W=240. At stride 8, acquired indices include 0,8,...,152,154 (21 acquired and 134 missing); at stride 16 they include 0,16,...,144,154 (11 acquired and 144 missing). A target at z=12 bounded by z=8 and z=16 has alpha=0.5, depth=12/154 and normalized gap=0.5. Context is computed only from those acquired slices. The trained VAE maps a 240x240 image to four 30x30 latent channels: 8x spatial downsampling per axis and 16x scalar-element reduction for FLAIR. Code path and results are supplied in sections 15-17.',
    15:results_text()+'\nHistorical evidence retained separately: ten-volume evaluation recorded diffusion PSNR_missing 27.785 dB at stride 8 and 25.260 dB at stride 16; these earlier results used different preprocessing and must not be pooled with the corrected experiment. Best stored VAE checkpoint is epoch 13, validation loss 0.0891816; diffusion checkpoint is epoch 48, noise loss 0.1177198. Those losses are not image-fidelity metrics.\n'+equations,
    16:'Python 3.10; measured environment PyTorch 2.9.1+cu128 on NVIDIA RTX 5060 Laptop GPU (approximately 8 GB VRAM). From the project directory, use the approved interpreter D:\\Kaggle competetion\\torch_env\\Scripts\\python.exe. Example commands: -m pytest tests -q; scripts/reconstruct_volume.py --vol_id 14 --stride 8 --seed 42 --steps 25; scripts/evaluate_all.py --vols 10 --strides 8 16; scripts/finetune_diffusion.py --init-checkpoint experiments/checkpoints/diffusion/best_model.pt --strides 4 8 16 --lr 2e-5 --time-limit-minutes 70 --output-dir experiments/review/new_run. scripts/review_experiment.py reproduces the complete bounded protocol. Resume uses an explicit --resume checkpoint. Source implements epsilon prediction and DDIM; DDPM inference and arbitrary encoder backbones are not implemented.',
    17:'Local data: 57,195 HDF5 slice files representing 369 volume IDs, each 155 axial planes at 240x240. Each image has four stored channels; current code assumes FLAIR/T1/T1ce/T2 ordering, which requires original HDF5 converter provenance confirmation. Three uint8 mask channels exist; sampled channels are mutually exclusive, so WT/TC/ET semantics are not assumed. Original BraTS is multi-institutional brain-tumor MRI with expert annotations (R6). Patient-volume split uses seed 42: 258 train, 55 validation, 56 test. These are internal partitions of the local training collection, not official challenge test data. Diffusion trains only on missing target planes. Fine-tuning uses equal-probability strides 4/8/16 and acquired-only intensity statistics; VAE stays frozen and encoder/U-Net receive MSE noise-prediction updates with AdamW and gradient clipping. Five fixed validation volumes select settings; ten held-out volumes provide final comparisons. No external data or third-party clinical study is used.'+training_note,
    18:'Requires inventor confirmation: earliest conception date, any public repository, uploaded slides, seminar/defense presentations, posters, abstracts, thesis deposits, demonstrations, publication dates, audiences, recipients and confidentiality conditions. Local presentation files exist; possession of a file does not establish that it was publicly disclosed. No declaration of no prior disclosure is made.',
    19:'Implemented research software with pretrained models, retrospective benchmark evidence and regression verification. Completed as an experimental prototype; diagnostic safety and clinical efficacy are not validated. Continued improvement depends on external-data testing, physical geometry provenance, lesion-specific evaluation and workflow studies.',
    20:'Proposed technical claim concepts for review, without an established novelty conclusion:\n1. A computer-implemented method receiving sparse registered MRI slices and acquired depth indices; calibrating intensity using acquired values; encoding bounding slices with shared weights; encoding normalized depth, relative offset and gap; conditioning latent denoising through spatial maps and attention tokens; decoding missing planes; assembling a dense volume while retaining acquired planes unchanged.\n2. The method of concept 1 using percentile normalization from the acquired stack and a frozen variational decoder.\n3. The method of concept 1 using Fourier coordinate embeddings, batched DDIM sampling and optional depth smoothing with acquired-plane restoration.\n4. A processing system or stored program configured to carry out the supported method.\nThese combinations require full prior-art review and confirmed inventorship. Improvements made during this review must be distinguished from the inventors’ original contributions and approved by them before adopting claim language.',
    21:'A preliminary technical search was performed on 8 October 2026 using arXiv, Google Patents and primary manufacturer/data sources. R1-R4 were examined for relevant disclosed mechanisms; R5 is a search lead with unavailable full text. Related concepts are already known; the specific combination and earliest disclosure date remain to be assessed. This search is not exhaustive, and patent-family claims/status were not independently confirmed in official registers. Comparison rows follow.',
    22:'A conclusion that an ordinarily skilled person would not arrive at this system cannot be supported from the current evidence. Latent diffusion, cross-attention, neighboring-slice synthesis, positional embeddings and intensity normalization are established techniques. The particular combination may merit examination, but organ choice or implementation detail alone does not establish inventive step. Requires inventor confirmation of the original technical contribution and dated development records.',
    23:'Demonstrated/configured settings: D=155; H=W=240; trained C=1; four-channel I/O tested without a trained joint checkpoint; acquired stride 2/4/8/16 in historical data and 8/16 in corrected evaluation; training strides 4/8/16; latent channels=4, spatial scale=8, latent factor=0.18215; context tokens=17, token dimension=256, spatial channels=8; encoder base=32; U-Net base=64 with multipliers 1/2/4; noise T=1000 with cosine beta schedule; epsilon prediction; DDIM 25 or 50 steps, eta=0; batch=8 for experiments, batch=2 for smoke verification; fine-tuning LR=2e-5, weight decay=1e-4, clipping norm=1, EMA decay=0.999; normalization percentiles 0.5/99.5, target [-1,1], epsilon=1e-7; optional smoothing sigma=0.5 along depth. API accepts positive strides and 1..T inference steps; mathematical acceptance is not evidence of quality across every value. Broader workable ranges require separate experiments.',
    24:'Potential companies, with interest unconfirmed:\n1. Siemens Healthcare GmbH, Henkestr. 127, 91052 Erlangen, Germany (official contact page, R7).\n2. Koninklijke Philips N.V., Philips Center, Prinses Irenestraat 59, 1077 WV Amsterdam, Netherlands (2025 filing address, R9).\nThese are prospective research/licensing contacts; no outreach or relationship is claimed.\nMarketing profile: research software for retrospective sparse brain MRI reconstruction using latent generation, documented acquired-data preservation and local GPU execution. Measured fidelity evidence is available; acquisition savings, clinical benefit, willingness to pay and product integration remain unvalidated. Commercial positioning must account for established latent interpolation work (R2) and scanner reconstruction products (R8).',
    25:'A credible market size for this specific software has not been established. Context: Philips reported approximately EUR 8.5 billion 2025 Diagnosis & Treatment sales and forecast mid-single-digit comparable sales CAGR for 2026-2028 (R11). This broad business segment includes multiple imaging and therapy products and is not a market-size or revenue estimate for this invention. Existing AI imaging products suggest an active category, but no quantified adoption study or customer demand is available for this prototype. Market sizing, licensing price, integration cost and clinical-economic savings require independent customer research.',
    26:'Requires inventor confirmation of the complete inventor list, legal name spellings, contribution descriptions, affiliations, addresses and ownership. The proposal lists Gudikandula Sai Charan chary (23BAI1536), Ayush chaurasia (23BAI1543), and Chahanaa Sheth (23BAI1517); these are candidate project contributors, not confirmed legal inventors. Dr. Sarita Kumari is named as project guide at SCOPE, VIT Chennai; supervision alone does not establish inventorship. Funding, institutional assignment and third-party contributions require confirmation.',
    27:'Potential operators: medical-imaging researchers, engineers and trained radiology professionals performing retrospective research. Intended future beneficiaries: patients whose MRI data are processed, after appropriate validation. Operator age: adults with relevant training; no validated patient age range is established. Expected research benefit: dense volume visualization and comparison of reconstruction methods, with acquired data preserved. No measured financial advantage or diagnostic benefit is claimed; local GPU use is demonstrated but acquisition cost, scan savings and licensing costs are unknown.',
    28:'Provisional TRL 4: integrated software prototype evaluated on retrospective data in a local laboratory/computing environment. Requires inventor/institution confirmation. No operational clinical environment, qualified clinical product or deployment evidence supports TRL 5-9. The selection below is an evidence-based draft assessment.',
    29:'Prepared from repository artifacts and a preliminary source search on 8 October 2026. Original templates, presentations, checkpoints and historical results remain preserved. Missing ground-truth planes are reference data only; normalization uses acquired slices. Research outputs can contain plausible but incorrect anatomy. Unverified physical geometry, HDF5 label provenance, external validity and public disclosure history are explicitly flagged. Inventors should confirm attribution, original contributions, administrative facts and final technical statements before this draft is used as their disclosure.'}

def append_answer(cell, text):
    # Preserve nested comparison/market tables and the outer template structure.
    for p in cell.paragraphs:
        if p.text.strip() in ('','.'):
            p._element.getparent().remove(p._element)
    for line in text.split('\n'):
        p=cell.add_paragraph(line)
        p.paragraph_format.space_after=Pt(5)
        for run in p.runs: run.font.size=Pt(10)

def build():
    diagrams()
    template=next(ROOT.glob('Invention DisclosureForm*.docx'))
    doc=Document(template)
    table=doc.tables[0]
    row_map={1:1,2:3,3:5,4:7,5:10,6:12,7:14,8:16,9:18,10:20,11:22,12:24,13:26,14:28,15:30,
        16:32,17:34,18:36,19:38,20:41,21:43,22:45,23:47,24:49,25:52,26:54,27:56,28:64,29:66}
    content=answers()
    for n,row in row_map.items(): append_answer(table.rows[row].cells[1],content[n])
    # Explicitly answer separate subquestion rows in the original form.
    supplemental={8:'Interpolation cannot recover anatomy unique to missing planes; learned priors can produce unsupported detail. This project has no measurement-consistency model, uncertainty calibration or clinical validation.',
        39:'The software prototype and retrospective fidelity experiments are complete; clinical efficacy and independent third-party validation are absent.',
        50:'Research software for retrospective sparse FLAIR MRI reconstruction. Supports acquired-data preservation, documented local GPU operation and measured image fidelity. No clinical, acquisition-time or economic advantage over established products is demonstrated.',
        57:'Adult research/engineering operators with relevant training. No validated patient age range is established.',
        58:'Potential research benefit: dense volume visualization and reconstruction-method comparison. Diagnostic benefit remains unvalidated.',
        59:'No measured cost advantage. Local GPU execution is demonstrated; acquisition savings, licensing price and workflow cost require separate evidence.'}
    for row,text in supplemental.items(): append_answer(table.rows[row].cells[1],text)
    for name in ('architecture_labeled','architecture_unlabeled','operation_labeled','operation_unlabeled'):
        table.rows[24].cells[1].add_paragraph().add_run().add_picture(str(DRAWINGS/(name+'.png')),width=Inches(5.5))
    # Populate comparison tables already present in the supplied template.
    comparison=[['1','R1: latent diffusion','Latent generation/cross-attention','MRI bounding-slice and depth context'],
        ['2','R2: CaLID','Sparse MRI latent interpolation','Brain FLAIR implementation; broad overlap'],
        ['3','R3: ISCS','Correlated diffusion noise','Independent seeded noise plus optional smoothing'],
        ['4','R4: US11410271B2','Binary missing-slice interpolation','Grayscale conditional latent generation']]
    for nested in table.rows[43].cells[1].tables:
        for i,values in enumerate(comparison,1):
            if i>=len(nested.rows): nested.add_row()
            for c,value in zip(nested.rows[i].cells,values): c.text=value
    for nested in table.rows[52].cells[1].tables:
        data=[['2025','Specific market size not established; contextual segment sales EUR 8.5bn (R11)','Imaging workflow and AI reconstruction','Deep Resolve (R8); broad Philips imaging portfolio'],
              ['2026–2028','Specific market forecast unavailable; Philips segment forecast mid-single-digit CAGR (R11)','Manufacturer investment; project demand unconfirmed','No commercial product from this project']]
        for i,values in enumerate(data,1):
            if i>=len(nested.rows): nested.add_row()
            for c,value in zip(nested.rows[i].cells,values): c.text=value
        for row in nested.rows[len(data)+1:]:
            for cell in row.cells: cell.text='Not established; requires further market research.'
    # Tick only the TRL4 column; retain the explanatory paragraph in the merged answer row.
    for i,c in enumerate(table.rows[62].cells):
        if c.text.strip()=='TRL 4':
            c.text='TRL 4 [selected, provisional]'
            break
    for row in table.rows:
        row.height=None
        row.height_rule=WD_ROW_HEIGHT_RULE.AUTO
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next=False
                p.paragraph_format.keep_together=False
    # Ensure comparison and market tables use readable fonts and flexible row heights.
    nested_seen=set()
    for cell in table._cells:
        for nested in cell.tables:
            if id(nested._tbl) in nested_seen: continue
            nested_seen.add(id(nested._tbl))
            for row in nested.rows:
                row.height=None
                row.height_rule=WD_ROW_HEIGHT_RULE.AUTO
                for c in row.cells:
                    for p in c.paragraphs:
                        for run in p.runs: run.font.size=Pt(10)
    doc.add_page_break()
    heading = doc.add_paragraph('Source register and evidence attachments')
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(14)
    doc.add_paragraph('Sources accessed 8 October 2026. R5 is a bibliographic lead only. Local results and checkpoint metadata are project evidence, not independent clinical validation.')
    for code,title,url in SOURCES: doc.add_paragraph(f'{code}. {title}\n{url}')
    doc.add_paragraph('Local evidence: PROJECT_REVIEW_PLAN.md; results/review/REVIEW_FINDINGS.md; PRIOR_ART.md; experiment_manifest.json; test_reference.csv; validation_reference.csv; validation_candidate.csv; selection.json; test_selected.csv; paired_comparison.csv; completion.json; experiments/review/finetune/manifest.json and training.csv; source scripts and regression tests.')
    target=ROOT/'Invention_Disclosure_Filled.docx'
    doc.save(target)
    (OUT/'disclosure_sections.json').write_text(json.dumps(content,indent=2,ensure_ascii=False),encoding='utf-8')
    (OUT/'sources.json').write_text(json.dumps(SOURCES,indent=2,ensure_ascii=False),encoding='utf-8')
    (OUT/'INVENTOR_CONFIRMATION.md').write_text('# Inventor confirmation required\n\n- Confirm legal inventors, name spellings, contributions and affiliations; proposal names are candidates only.\n- Confirm ownership, institutional assignments, funding and external contributions.\n- Provide earliest conception and public disclosure dates, recipients and confidentiality conditions.\n- Confirm HDF5 modality ordering, mask label mapping and conversion provenance.\n- Confirm which technical features were original contributions and distinguish later review corrections.\n- Review preliminary prior-art comparisons, proposed claim concepts and provisional TRL 4.\n- Confirm any independent validation, commercial contact, market estimate or savings evidence; none has been established.\n',encoding='utf-8')
    print(target)

if __name__=='__main__': build()
