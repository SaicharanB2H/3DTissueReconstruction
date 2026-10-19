"""Record the completed structural and native Word visual review."""
import sys,json,hashlib,re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.review_tools'))
from docx import Document
from finalize_disclosure import ROWS,image_hashes
OUT=ROOT/'results/review/disclosure_final'
docpath=ROOT/'Invention_Disclosure_Final.docx'
doc=Document(docpath);source=Document(ROOT/'Invention_Disclosure_Revised_v2.docx')
answers=json.loads((OUT/'disclosure_sections.json').read_text(encoding='utf-8'))
missing=[]
for n,row in ROWS.items():
    cell=doc.tables[0].rows[row].cells[0 if n==28 else 1]
    if answers[str(n)].strip() not in cell.text.strip():missing.append(n)
assert not missing,f'Unpopulated or altered sections: {missing}'
assert sorted(image_hashes(doc))==sorted(image_hashes(source))
sources=json.loads((ROOT/'results/review/disclosure_revision/sources.json').read_text(encoding='utf-8'))
for entry in sources:
    if entry['id']=='R9':entry['publication_date_or_version']=entry['publication_date_or_version'].replace('2026 SEC filing','SEC filing dated 19 February 2026')
sources.append({'id':'R20','authors_or_organization':'Jiaming Song; Chenlin Meng; Stefano Ermon',
    'title':'Denoising Diffusion Implicit Models','publication_date_or_version':'arXiv:2010.02502; v1 6 October 2020; v4 5 October 2022; ICLR 2021',
    'url':'https://arxiv.org/abs/2010.02502','accessed':'2026-10-08'})
(OUT/'sources.json').write_text(json.dumps(sources,indent=2,ensure_ascii=False),encoding='utf-8')
pages=json.loads((OUT/'native_page_text.json').read_text(encoding='utf-8'))
assert len(pages)==18
fulltext='\n'.join(x['text'] for x in pages)
assert all(f'R{i}.' in fulltext for i in range(1,21))
assert all(f'Figure {i}.' in pages[i+2]['text'] for i in range(1,5))
assert 'Provide available data or information on market potential' in pages[13]['text']
verification=json.loads((OUT/'verification.json').read_text(encoding='utf-8'))
digest=hashlib.sha256(docpath.read_bytes()).hexdigest()
assert digest==verification['final_sha256'],'Document changed after rendering'
verification.update({'native_render':'success; Microsoft Word UI PDF export, all 18 pages visually inspected',
    'native_pdf_sha256':hashlib.sha256((OUT/'native_word_final.pdf').read_bytes()).hexdigest(),
    'native_page_count':18,'visual_reviewed_pages':list(range(1,19)),
    'render_dpi':144,'Q25_verified_page':14,'figure_pages':{'1':4,'2':5,'3':6,'4':7},
    'TRL_verified_page':15,'prior_art_table_page':12,'parameter_table_page':13,
    'all_sections_populated':True,'source_register_entries':20,
    'visual_issues_remaining':[], 'finalized_date':'2026-10-09',
    'diagram_sha256_in_final_order':image_hashes(doc),
    'layout_note':'Original form retained; drawing/evidence tables use inline grouping rows to prevent splitting in Word compatibility mode.'})
(OUT/'verification.json').write_text(json.dumps(verification,indent=2),encoding='utf-8')
(OUT/'native_render_status.json').write_text(json.dumps({'status':'success','renderer':'Microsoft Word native Save As PDF through UI',
    'document_sha256':digest,'pdf':'native_word_final.pdf','pages':18,'visual_review':'all pages at original PNG resolution',
    'earlier_hidden_COM_export':'stalled; superseded by successful native UI export',
    'fallback_renderer':'Official LibreOffice 26.2.6 MSI verified by official SHA-256 and valid Authenticode signature; administrative extraction only; PDF export succeeded but native Word export is the acceptance evidence'},indent=2),encoding='utf-8')
print(json.dumps({'final_sha256':digest,'sections':29,'pages':18,'diagrams_unchanged':True,'sources':20},indent=2))
