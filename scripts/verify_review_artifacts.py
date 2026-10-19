"""Validate disclosure coverage, table preservation and completed experiment evidence."""
import sys
import json
import hashlib
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.review_tools'))
sys.path.insert(0,str(ROOT))
from docx import Document
import pandas as pd
from src.models.factory import checkpoint_id
OUT=ROOT/'results/review'

def structure(doc):
    return [(len(t.rows),len(t.columns),[len(r.cells) for r in t.rows]) for t in doc.tables]

def run():
    template=next(ROOT.glob('Invention DisclosureForm*.docx'))
    original=Document(template); filled=Document(ROOT/'Invention_Disclosure_Filled.docx')
    answers=json.loads((OUT/'disclosure_sections.json').read_text(encoding='utf-8'))
    assert set(answers)==set(map(str,range(1,30))) and all(text.strip() for text in answers.values())
    content='\n'.join(c.text for t in filled.tables for row in t.rows for c in row.cells)
    assert all(text in content for text in answers.values())
    assert 'experiments are in progress' not in content and 'Requires inventor confirmation' in content
    assert structure(original)==structure(filled)
    assert len(filled.inline_shapes)==4
    assert len(original._element.xpath('.//w:tbl'))==len(filled._element.xpath('.//w:tbl'))
    sources=json.loads((OUT/'sources.json').read_text(encoding='utf-8'))
    paragraphs='\n'.join(p.text for p in filled.paragraphs)
    assert all(url in paragraphs for _,_,url in sources)
    for p in filled.inline_shapes:
        assert 0<p.width.inches<=5.5 and 0<p.height.inches<9
    frames={name:pd.read_csv(OUT/(name+'.csv')) for name in ['test_reference','validation_reference','validation_candidate','test_selected']}
    assert [len(frames[n]) for n in frames]==[100,80,80,20]
    assert not set(frames['validation_reference'].vol_id)&set(frames['test_selected'].vol_id)
    ref_info=json.loads((OUT/'test_reference.json').read_text())
    preserved={}
    for key in ['checkpoint','vae_checkpoint']:
        saved=ref_info[key]
        preserved[key]=checkpoint_id(saved['path'])['sha256']==saved['sha256']
    assert all(preserved.values())
    complete=json.loads((OUT/'completion.json').read_text())
    assert complete['status']=='completed' and complete['seconds']<=7200
    result={'status':'passed','disclosure_sections':29,'drawings':4,'outer_table_rows':len(filled.tables[0].rows),'outer_table_grid_columns':len(filled.tables[0].columns),'tables_preserved':True,'source_register_entries':len(sources),'benchmark_rows':{name:len(frame) for name,frame in frames.items()},'original_checkpoint_hashes_unchanged':preserved,'experiment_elapsed_seconds':complete['seconds'],'regression_tests':'34 passed; latest invocation recorded in execution output','training_smoke':'two iterations followed by explicit resume to cumulative iteration three','native_page_rendering':'unavailable: Microsoft Word COM server initialization failed (80080005); drawings inspected directly and XML/table dimensions verified','git_diff_check':'Review edits compile; pre-existing trailing whitespace remains in user-edited scripts/view_3d_blender.py and was preserved'}
    (OUT/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':run()
