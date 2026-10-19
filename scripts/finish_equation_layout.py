from pathlib import Path
import sys, json, hashlib, zipfile, shutil
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.review_tools'))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=ROOT/'results/review/disclosure_equations'
SOURCE=ROOT/'Invention_Disclosure_Final_Equations.docx'
FINAL=ROOT/'Invention_Disclosure_Word_Equations_Final.docx'
shutil.copy2(SOURCE,OUT/'Equation_Copy_Before_Final_Layout.docx')
d=Document(SOURCE)
# Preserve the form's original compatibility layout; Word's Save As upgrade
# otherwise permits a nested drawing row to split between caption and picture.
for item in d.settings._element.xpath('.//w:compatSetting'):
    if item.get(qn('w:name'))=='compatibilityMode': item.set(qn('w:val'),'14')
# Keep the clean architecture caption on its drawing page in modern Word too.
p=d.tables[0].rows[24].cells[1].tables[2].cell(0,0).paragraphs[0]
p.paragraph_format.page_break_before=True
p.paragraph_format.keep_with_next=True
# Put PSNR on a separate line so Word does not break its product mid-equation.
for p in d._element.xpath('.//w:p'):
    if ''.join(p.xpath('.//w:t/text()')).startswith('Reconstruction '):
        for m in p.xpath('./m:oMath'):
            if ''.join(m.xpath('.//m:t/text()',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})).startswith('PSNR'):
                r=OxmlElement('w:r');r.append(OxmlElement('w:br'));m.addprevious(r)
# Retain editable OMML and approved figure media after the local layout adjustments.
d.save(FINAL)
def media(path):
    with zipfile.ZipFile(path) as z:return {x:hashlib.sha256(z.read(x)).hexdigest() for x in z.namelist() if x.startswith('word/media/')}
assert sorted(media(FINAL).values())==sorted(media(OUT/'Invention_Disclosure_Before_Equation_Formatting.docx').values())
audit=json.loads((OUT/'verification.json').read_text())
audit.update({'final_document':str(FINAL),'docx_sha256':hashlib.sha256(FINAL.read_bytes()).hexdigest(),'native_equations':len(d._element.xpath('.//m:oMath')),'latest_open_copy_preserved':True})
(OUT/'verification.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print(FINAL)
