from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.review_tools'))
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
path=ROOT/'BCSE497J_Project_Report_23BAI1536.docx'
d=Document(path)
t=d.tables[12]
for i,w in enumerate([1.9,3.25,4.75,4.75]):
    t.columns[i].width=Cm(w)
    for r in t.rows:r.cells[i].width=Cm(w)
inside=False
for p in d.paragraphs:
    if p.text=='REFERENCES':inside=True;continue
    if p.text.startswith('APPENDIX 1'):inside=False
    if inside:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
d.save(path)
print('Corrected summary column widths and APA reference alignment')
