"""Record project inventory, document contents and dataset structure without altering originals."""
import sys
import json
import hashlib
import re
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.review_tools'))
sys.path.insert(0, str(ROOT))
import h5py
import numpy as np
import pandas as pd
import pdfplumber
from docx import Document
from src.models.factory import read_configs
from src.data.dataset import configured_patient_splits

OUT = ROOT / 'results/review'

def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1048576), b''): h.update(chunk)
    return h.hexdigest()

def run():
    extracted = OUT / 'document_extraction'
    extracted.mkdir(parents=True, exist_ok=True)
    inventory = []
    paths = list(ROOT.glob('*')) + list((ROOT/'src').rglob('*')) + list((ROOT/'scripts').rglob('*')) + list((ROOT/'tests').rglob('*')) + list((ROOT/'configs').rglob('*')) + list((ROOT/'PPT').rglob('*'))
    for p in sorted(set(paths)):
        if not p.is_file() or '__pycache__' in p.parts or p.name.startswith('Invention_Disclosure_Filled'): continue
        record = {'path':str(p.relative_to(ROOT)), 'bytes':p.stat().st_size, 'sha256':digest(p)}
        if p.suffix.lower() in ('.py','.md','.yaml','.txt','.csv','.mtl'):
            text = p.read_text(encoding='utf-8',errors='replace')
            record['lines'] = len(text.splitlines())
        elif p.suffix.lower()=='.pptx':
            chunks=[]
            with zipfile.ZipFile(p) as z:
                names=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)], key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
                for n in names:
                    chunks.append(n+'\n'+'\n'.join(e.text or '' for e in ET.fromstring(z.read(n)).iter() if e.tag.endswith('}t')))
            (extracted/(p.name+'.txt')).write_text('\n\n'.join(chunks),encoding='utf-8')
            record['slides']=len(chunks)
        elif p.suffix.lower()=='.pdf':
            with pdfplumber.open(p) as pdf:
                (extracted/(p.name+'.txt')).write_text('\n\n'.join(f'PAGE {i+1}\n'+(page.extract_text() or '') for i,page in enumerate(pdf.pages)),encoding='utf-8')
                record['pages']=len(pdf.pages)
        elif p.suffix.lower()=='.docx':
            doc=Document(p)
            chunks=[x.text for x in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(' | '.join(dict.fromkeys(c.text for c in row.cells)))
            (extracted/(p.name+'.txt')).write_text('\n'.join(chunks),encoding='utf-8')
            record['tables']=len(doc.tables)
        elif p.suffix.lower()=='.zip':
            with zipfile.ZipFile(p) as z: record['members']=z.namelist()
        inventory.append(record)
    (OUT/'project_inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
    cfg=read_configs()[0]
    data=ROOT/cfg['data']['data_dir']
    found=defaultdict(set)
    for p in data.glob('volume_*_slice_*.h5'):
        m=re.fullmatch(r'volume_(\d+)_slice_(\d+)\.h5',p.name)
        found[int(m.group(1))].add(int(m.group(2)))
    incomplete={str(v):sorted(set(range(155))-indices) for v,indices in found.items() if indices != set(range(155))}
    meta=pd.read_csv(data/'meta_data.csv')
    samples=[]
    for v in [1,14,100,369]:
        if v not in found: continue
        for depth in [0,75,154]:
            p=data/f'volume_{v}_slice_{depth}.h5'
            with h5py.File(p) as f:
                arr=f['image'][()]; mask=f['mask'][()]
                samples.append({'volume':v,'slice':depth,'image_shape':list(arr.shape),'image_dtype':str(arr.dtype),'finite':bool(np.isfinite(arr).all()),'mask_shape':list(mask.shape),'mask_dtype':str(mask.dtype),'mask_channel_sums':mask.sum(axis=(0,1)).tolist(),'mask_overlapping_pixels':int((mask.sum(axis=-1)>1).sum()),'attributes':list(f.attrs.keys())})
    splits=configured_patient_splits(cfg,42)
    audit={'h5_files':sum(map(len,found.values())),'volumes':len(found),'incomplete_volumes':incomplete,'metadata_rows':len(meta),'metadata_duplicate_volume_slice_pairs':int(meta.duplicated(['volume','slice']).sum()),'metadata_ids_match_files':set(meta.volume.unique())==set(found),'patient_splits':dict(zip(['train','validation','test'],splits)),'split_disjoint':all(not set(splits[a])&set(splits[b]) for a,b in [(0,1),(0,2),(1,2)]),'sample_inspection':samples,'limits':'All filenames/indices and metadata inspected; HDF5 payload inspection is sampled, not an exhaustive voxel integrity scan. Converter provenance and physical geometry are unavailable.'}
    (OUT/'dataset_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
    print(json.dumps({'inventory_files':len(inventory),'dataset_files':audit['h5_files'],'volumes':audit['volumes'],'incomplete':len(incomplete),'disjoint':audit['split_disjoint']}))

if __name__=='__main__': run()
